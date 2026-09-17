r"""Ministry of Youth and Sports (GSB, Spor Hizmetleri Genel Müdürlüğü) statistics by province.

shgm.gsb.gov.tr/Sayfalar/175/105/Istatistikler links a handful of workbooks, kept in
`C:\veri-ham\gsb`. Two are by province:

* "İllere Göre Kulüp Sayıları 2025": sports clubs, one year only;
* "Sportif Yetenek Taraması ve Spora Yönlendirme Programı 2022-2025": pupils screened (with and
  without disability), found suited to sport, and directed to a branch, one sheet per year and
  stage. The directing stage runs a season behind (the 2024 sheet says so), and 2025 has no
  directing sheet with data yet.

Athletes, coaches and referees are published for Türkiye by federation only and are not loaded.

Checks: 81 provinces per sheet; they add up to the sheet's TOPLAM / Genel Toplam row.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import polars as pl

from ..config import RAW
from .kgm import fold, province_id

FOLDER = RAW / "gsb" if (RAW / "gsb").exists() else Path("C:/veri-ham/gsb")
CLUBS = "İllere_Göre_Kulüp_Sayıları_2025.xlsx"
TALENT = "Sportif_Yetenek_Taraması_ve_Spora_Yönlendirme_Programı_2022-2025.xlsx"
#: folded sheet-name fragment -> (indicator, {column header fragment: dims})
TALENT_SHEETS = {
    "geneltarama": (
        "gsb_talent_screened",
        {
            "engelsiz": "disability=without",
            "engelli": "disability=with",
            "toplam": "disability=total",
        },
    ),
    "sporayatkin": ("gsb_talent_suited", {"": ""}),
    "yonlendiril": ("gsb_talent_directed", {"": ""}),
}
Row = tuple[str, str, str, int]  # indicator, area, dims, year


def rows_of(path: Path, sheet=None) -> list[list]:
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = book[sheet] if sheet else book.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def province_rows(
    rows: list[list], columns: dict[int, str], what: str
) -> dict[tuple[str, str], float]:
    """{(area, dims): value}; the provinces must add up to the TOPLAM row."""
    out: dict[tuple[str, str], float] = {}
    printed: dict[str, float] = {}
    for r in rows:
        name = next((c for c in r if isinstance(c, str) and c.strip()), None)
        if name is None:
            continue
        key = fold(name)
        if key in ("toplam", "geneltoplam"):
            printed = {dims: float(r[j] or 0) for j, dims in columns.items()}
            continue
        try:
            area = province_id(name.strip())
        except KeyError:
            continue
        for j, dims in columns.items():
            value = r[j]
            if not isinstance(value, (int, float)):
                continue
            if (area, dims) in out:
                raise ValueError(f"GSB {what}: {name} iki kez")
            out[(area, dims)] = float(value)
    areas = {a for a, _ in out}
    if len(areas) != 81:
        raise ValueError(f"GSB {what}: {len(areas)} il")
    if not printed:
        raise ValueError(f"GSB {what}: TOPLAM satırı yok")
    for dims, total in printed.items():
        read = sum(v for (a, d), v in out.items() if d == dims)
        if abs(read - total) > 0.5:
            raise ValueError(
                f"GSB {what} {dims}: iller {read:,.0f}, TOPLAM {total:,.0f}"
            )
    return out


def load_all() -> dict[Row, float]:
    out: dict[Row, float] = {}
    rows = rows_of(FOLDER / CLUBS)
    header = next(
        i
        for i, r in enumerate(rows)
        if any(isinstance(c, str) and fold(c) == "il" for c in r)
    )
    column = next(
        j
        for j, c in enumerate(rows[header])
        if isinstance(c, str) and "kulub" in fold(c)
    )
    for (area, dims), value in province_rows(
        rows[header + 1 :], {column: ""}, "kulüp"
    ).items():
        out[("gsb_sports_clubs", area, dims, 2025)] = value

    import openpyxl

    book = openpyxl.load_workbook(FOLDER / TALENT, read_only=True, data_only=True)
    for sheet in book.sheetnames:
        year = int(re.match(r"(\d{4})", sheet).group(1))
        key = fold(sheet)
        indicator, wanted = next(v for k, v in TALENT_SHEETS.items() if k in key)
        rows = rows_of(FOLDER / TALENT, sheet)
        head = next(
            (
                i
                for i, r in enumerate(rows)
                if any(isinstance(c, str) and fold(c) == "il" for c in r)
            ),
            None,
        )
        body = rows[head + 1 :] if head is not None else []
        if not any(isinstance(c, (int, float)) and c for r in body for c in r[2:3]):
            continue  # 2025 directing sheet: no data yet
        columns = {}
        for j, cell in enumerate(rows[head]):
            if j < 2 or not isinstance(cell, str):
                continue
            folded = fold(cell)
            for fragment, dims in wanted.items():
                if fragment in folded and dims not in columns.values():
                    columns[j] = dims
                    break
        for (area, dims), value in province_rows(body, columns, sheet).items():
            out[(indicator, area, dims, year)] = value
    return out


_CACHE: dict[Row, float] = {}
UNITS = {
    "gsb_sports_clubs": "item",
    "gsb_talent_screened": "person",
    "gsb_talent_suited": "person",
    "gsb_talent_directed": "person",
}


class Gsb:
    source_id = "gsb"
    indicator_id = ""

    def fetch(self) -> Path:
        return FOLDER

    def parse(self, raw: Path) -> pl.DataFrame:
        if not _CACHE:
            _CACHE.update(load_all())
        records = [
            {
                "area_id": area,
                "area_level": "province",
                "period_start": dt.date(year, 1, 1),
                "dims": dims,
                "value": value,
            }
            for (indicator, area, dims, year), value in _CACHE.items()
            if indicator == self.indicator_id
        ]
        return pl.DataFrame(
            records, schema_overrides={"value": pl.Float64}
        ).with_columns(
            pl.lit(self.indicator_id).alias("indicator_id"),
            pl.lit("annual").alias("frequency"),
            pl.lit(UNITS[self.indicator_id]).alias("unit"),
            pl.lit("measured").alias("quality_flag"),
            pl.lit("2026-09").alias("vintage"),
            pl.lit(self.source_id).alias("source_id"),
            pl.lit(dt.date(2026, 9, 17)).alias("retrieved_at"),
        )


GSB_ADAPTERS = {
    indicator: type(f"Gsb_{indicator}", (Gsb,), {"indicator_id": indicator})
    for indicator in UNITS
}
