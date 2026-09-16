r"""DHMİ airport statistics: aircraft, passengers, freight and cargo by airport, monthly.

The workbooks (`C:\veri-ham\dhmi\<id>.xlsx`, downloaded by `scripts/fetch_dhmi.py`) each hold
five sheets — all aircraft, commercial aircraft, passengers, freight (baggage + cargo + mail,
tonnes) and cargo (tonnes) — and each sheet is a year-to-date table: one row per airport, and
two column groups, the same month of the previous year and of the report's own year, each
split into domestic, international and total.

So a month is printed by two years' worth of files: as the current year in its own file and as
the comparison column a year later. That is the check — and it is also how 2011 survives,
whose own files are not on the site any more.

The figures are cumulative, so the monthly value is the difference from the month before
(January is the cumulative figure itself). Where the previous month is missing, the month is
dropped rather than guessed.

Sheets are recognised by their title cell, because the sheet names change from year to year
("TÜM UÇAK", "TUM_UCAK", "TÜM", "Sayfa1"). The airport rows carry the province in their name;
`AIRPORTS` maps the exceptions. Airports run by others are marked "(*)" in the workbook and
left out of DHMİ's own total, so the printed total row is not the sum of the airports and is
not used either.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import polars as pl

from ..config import RAW
from .kgm import fold, province_id

FOLDER = Path("C:/veri-ham/dhmi") if not (RAW / "dhmi").exists() else RAW / "dhmi"
MONTHS = {
    "OCAK": 1,
    "ŞUBAT": 2,
    "MART": 3,
    "NİSAN": 4,
    "MAYIS": 5,
    "HAZİRAN": 6,
    "TEMMUZ": 7,
    "AĞUSTOS": 8,
    "EYLÜL": 9,
    "EKİM": 10,
    "KASIM": 11,
    "ARALIK": 12,
}
PERIOD = re.compile(r"(20\d\d)\s*(?:YILI)?\s*([A-ZÇĞİÖŞÜ]+)\s*SONU")
# sheet title -> (indicator, dims prefix)
MEASURES = {
    r"tumucaktrafigi": ("dhmi_air_traffic", "aircraft_type=all"),
    r"ticariucaktrafigi": ("dhmi_air_traffic", "aircraft_type=commercial"),
    r"yolcutrafigi": ("dhmi_air_passengers", ""),
    r"yuktrafigi": ("dhmi_air_freight", "freight_type=freight"),
    r"kargotrafigi": ("dhmi_air_freight", "freight_type=cargo"),
}
UNITS = {
    "dhmi_air_traffic": "item",
    "dhmi_air_passengers": "person",
    "dhmi_air_freight": "tonne",
}
ROUTES = ("domestic", "international")
#: canonical airport code -> (province, folded labels the workbooks use over the years)
AIRPORTS = {
    "adana": ("Adana", ("adana",)),
    "adiyaman": ("Adıyaman", ("adiyaman",)),
    "agri": ("Ağrı", ("agri", "agriahmedihani")),
    "gazipasa": ("Antalya", ("alanyagazipasa", "antalyagazipasa", "gazipasaalanya")),
    "merzifon": ("Amasya", ("amasyamerzifon", "amerzifon", "merzifon")),
    "esenboga": ("Ankara", ("ankaraesenboga", "esenboga")),
    "antalya": ("Antalya", ("antalya",)),
    "ataturk": ("İstanbul", ("ataturk", "istanbulataturk")),
    "cildir": ("Aydın", ("aydincildir",)),
    "balikesir": ("Balıkesir", ("balikesir", "balikesirmerkez")),
    "kocaseyit": ("Balıkesir", ("balikesirkocaseyit", "balikesirkorfez")),
    "batman": ("Batman", ("batman",)),
    "bingol": ("Bingöl", ("bingol",)),
    "yenisehir": ("Bursa", ("bursayenisehir",)),
    "canakkale": ("Çanakkale", ("canakkale",)),
    "gokceada": ("Çanakkale", ("canakkalegokceada", "gokceada")),
    "cardak": ("Denizli", ("cardak", "denizlicardak")),
    "cukurova": ("Mersin", ("cukurova",)),  # Çukurova airport is in Tarsus, Mersin
    "dalaman": ("Muğla", ("dalaman", "mugladalaman")),
    "diyarbakir": ("Diyarbakır", ("diyarbakir",)),
    "elazig": ("Elazığ", ("elazig",)),
    "erzincan": ("Erzincan", ("erzincan", "erzincanyildirimakbulut")),
    "erzurum": ("Erzurum", ("erzurum",)),
    "hasanpolatkan": (
        "Eskişehir",
        (
            "eskisehiranadolu",
            "eskisehiranadoluu",
            "eskisehiranadoluun",
            "eskisehirhasanpolatkan",
        ),
    ),
    "gaziantep": ("Gaziantep", ("gaziantep",)),
    "hakkari": ("Hakkari", ("hakkariyuksekovase", "hakkariyuksekovaselahaddineyyubi")),
    "hatay": ("Hatay", ("hatay",)),
    "igdir": ("Iğdır", ("igdir", "igdirsehitbulentaydin")),
    "suleymandemirel": ("Isparta", ("ispartasuleymandemirel", "sdemirel")),
    "istanbul": ("İstanbul", ("istanbul",)),
    "sabihagokcen": ("İstanbul", ("istanbulsabihagokcen", "istsabihagokcen")),
    "adnanmenderes": ("İzmir", ("izmiradnanmenderes", "amenderes")),
    "kahramanmaras": ("Kahramanmaraş", ("kahramanmaras", "kmaras")),
    "kapadokya": ("Nevşehir", ("kapadokya", "nevsehirkap", "nevsehirkapadokya")),
    "kars": ("Kars", ("kars", "karsharakani")),
    "kastamonu": ("Kastamonu", ("kastamonu",)),
    "kayseri": ("Kayseri", ("kayseri",)),
    "cengiztopel": ("Kocaeli", ("kocaelicengiztopel",)),
    "konya": ("Konya", ("konya",)),
    "malatya": ("Malatya", ("malatya",)),
    "mardin": ("Mardin", ("mardin", "mardinprofdrazizsancar")),
    "milasbodrum": ("Muğla", ("milasbodrum", "muglamilasbodrum")),
    "mus": ("Muş", ("mus", "mussultanalparslan")),
    "ordugiresun": ("Ordu", ("ordugiresun",)),  # the airport is in Gülyalı, Ordu
    "rizeartvin": ("Rize", ("rizeartvin",)),  # Pazar, Rize
    "carsamba": ("Samsun", ("samsuncarsamba", "scarsamba")),
    "gap": ("Şanlıurfa", ("sanliurfagap", "surfagap")),
    "siirt": ("Siirt", ("siirt",)),
    "sinop": ("Sinop", ("sinop",)),
    "sirnak": ("Şırnak", ("sirnakserafettinelci",)),
    "sivas": ("Sivas", ("sivas", "sivasnuridemirag")),
    "corlu": ("Tekirdağ", ("tekirdagcorlu", "tekirdagcorluataturk")),
    "tokat": ("Tokat", ("tokat",)),
    "trabzon": ("Trabzon", ("trabzon",)),
    "usak": ("Uşak", ("usak",)),
    "van": ("Van", ("vanferitmelen", "vanfmelen")),
    "zafer": ("Kütahya", ("zafer", "zaferbolgesel")),  # Altıntaş, Kütahya
    "caycuma": ("Zonguldak", ("zonguldakcaycuma",)),
}
BY_LABEL = {
    label: (code, province)
    for code, (province, labels) in AIRPORTS.items()
    for label in labels
}
#: rows that are not airports: the aggregates, the sheet titles and the footnotes
SKIP = re.compile(
    r"^(havaalanlari|havalimanlari|dhmi|diger|turkiye|overflight|toplam|genel"
    r"|tumucak|ticariucak|yolcutrafigi|yuktrafigi|kargotrafigi|revize|ekimayi|yukverileri|isaretli)"
)
Key = tuple[str, str, str, int, int]  # indicator, dims, airport code, year, month
ROW_GAPS: list[
    tuple
] = []  # cells where domestic + international misses the printed total


def airport_key(label: str) -> tuple[str, str] | None:
    """(airport code, province) for an airport row; None for aggregates and notes."""
    key = fold(re.sub(r"\(.*?\)|\*+", " ", label).split("(")[0])
    if not key or SKIP.match(key) or len(key) > 34:  # a footnote, not an airport
        return None
    found = BY_LABEL.get(key)
    if found is None:
        raise ValueError(f"DHMİ: tanınmayan havalimanı {label!r} ({key})")
    return found


def read_book(path: Path) -> dict[Key, float]:
    """Cumulative figures of one workbook, from both of its column groups."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[Key, float] = {}
    for name in book.sheetnames:
        rows = [row for row in book[name].iter_rows(values_only=True)]
        title = fold(" ".join(str(c) for c in rows[0][:3] if c)) if rows else ""
        measure = next(
            (
                (i, d)
                for pattern, (i, d) in MEASURES.items()
                if re.match(pattern, title)
            ),
            None,
        )
        if measure is None or len(rows) < 4:
            continue
        indicator, prefix = measure
        header = " ".join(str(c) for c in rows[1] if c)
        periods = PERIOD.findall(header.upper())
        if len(periods) != 2:
            continue
        groups = [
            (int(year), MONTHS[month], 1 + 3 * k)  # first data column of the group
            for k, (year, month) in enumerate(periods)
            if month in MONTHS
        ]
        for row in rows[3:]:
            label = str(row[0]) if row[0] else ""
            airport = airport_key(label)
            if airport is None:
                continue
            code, _province = airport
            for year, month, first in groups:
                cells = row[first : first + 3]
                if not all(isinstance(c, (int, float)) for c in cells):
                    continue
                domestic, international, total = (float(c) for c in cells)
                gap = abs(domestic + international - total)
                if gap > max(1.0, abs(total) * 1e-6):
                    # The printed "Toplam" column is not always the sum of its two parts —
                    # DHMİ rounds the tonnes, and some workbooks are simply wrong (2024-5
                    # İstanbul cargo misses by 12%). Only the two parts are loaded; the gaps
                    # are collected for the notes, and the reading is checked instead against
                    # the other workbooks that print the same month.
                    ROW_GAPS.append(
                        (path.name, name, label, domestic, international, total)
                    )
                for route, value in zip(ROUTES, (domestic, international), strict=True):
                    dims = ";".join(filter(None, (prefix, f"route={route}")))
                    out[(indicator, dims, code, year, month)] = value
    book.close()
    return out


def cumulative() -> tuple[dict[Key, float], list[tuple]]:
    """Every workbook's figures merged, and the disagreements between two workbooks."""
    paths = sorted(FOLDER.glob("*.xlsx"), key=lambda p: int(p.stem))
    if len(paths) < 150:
        raise FileNotFoundError(f"{FOLDER}: {len(paths)} dosya (scripts/fetch_dhmi.py)")
    out: dict[Key, float] = {}
    conflicts: list[tuple] = []
    for path in paths:
        for key, value in read_book(path).items():
            seen = out.get(key)
            if seen is not None and abs(seen - value) > max(1.0, abs(value) * 0.005):
                conflicts.append((path.name, key, seen, value))
            out[key] = value  # the newest workbook wins
    return out, conflicts


NEGATIVE: list[tuple] = []  # months that came out negative after a revision


def monthly() -> dict[tuple[str, str, str, dt.date], float]:
    """Cumulative differenced into months.

    A month whose predecessor is missing is dropped rather than guessed. A difference can come
    out negative when the two cumulative figures are from workbooks published either side of a
    revision (88 of 90,510 in 2026-09, nearly all tonnes of freight): those are dropped too and
    collected in ``NEGATIVE``.
    """
    figures, _conflicts = cumulative()
    out: dict[tuple[str, str, str, dt.date], float] = {}
    for (indicator, dims, code, year, month), value in figures.items():
        if month == 1:
            flow = value
        else:
            previous = figures.get((indicator, dims, code, year, month - 1))
            if previous is None:
                continue
            flow = value - previous
        if flow < 0:
            NEGATIVE.append((indicator, dims, code, year, month, flow))
            continue
        out[(indicator, dims, code, dt.date(year, month, 1))] = flow
    return out


class Dhmi:
    source_id = "dhmi"
    indicator_id = ""

    def fetch(self) -> Path:
        return FOLDER

    def parse(self, raw: Path) -> pl.DataFrame:
        records = []
        for (indicator, dims, code, start), value in monthly().items():
            if indicator != self.indicator_id:
                continue
            records.append(
                {
                    "area_id": province_id(AIRPORTS[code][0]),
                    "period_start": start,
                    "dims": f"airport={code};{dims}" if dims else f"airport={code}",
                    "value": value,
                }
            )
        return pl.DataFrame(
            records, schema_overrides={"value": pl.Float64}
        ).with_columns(
            pl.lit("province").alias("area_level"),
            pl.lit(self.indicator_id).alias("indicator_id"),
            pl.lit("monthly").alias("frequency"),
            pl.lit(UNITS[self.indicator_id]).alias("unit"),
            pl.lit("measured").alias("quality_flag"),
            pl.lit("2026-09").alias("vintage"),
            pl.lit("dhmi").alias("source_id"),
            pl.lit(dt.date(2026, 9, 16)).alias("retrieved_at"),
        )


DHMI_ADAPTERS = {
    indicator: type(f"Dhmi_{indicator}", (Dhmi,), {"indicator_id": indicator})
    for indicator in UNITS
}
