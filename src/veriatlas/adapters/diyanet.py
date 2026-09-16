r"""Diyanet statistics: mosques, personnel, Qur'an courses, pilgrims, converts, budget.

`scripts/fetch_diyanet.py` keeps the eleven workbooks of
stratejigelistirme.diyanet.gov.tr/sayfa/57/istatistikler in `C:\veri-ham\diyanet`.

Two shapes are read:

* Türkiye series — mosques and personnel 2013-2023 (tables 1.1, 2.1), the 2023 breakdowns of
  pilgrims by age and sex (5.1), converts (4.1) and the 2022-2023 budget (6.1).
* Provincial 2023 tables — mosques (2.2), personnel (1.3), Qur'an courses and their
  participants (3.3). These are printed by level-3 statistical region, and every level-3
  region is a single province, so they are loaded at province level; 2.2 prints the provinces
  in two columns side by side, so codes are looked for in every cell rather than in column A.

The personnel figures differ between the two tables on purpose: 1.1 counts the whole
Presidency (140,185 in 2023), 1.3 only the provincial and district muftiships (136,384) —
headquarters, Religious Affairs High Council and the like are outside it.

Checks that run on every load: the provincial figures must sum to the Türkiye row printed in
the same table; personnel by tenure (1.1) must sum to the total; and the graduates of the
Qur'an courses must come to the same number whether counted by sex, by level of education or
by age group.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import polars as pl

from ..config import RAW
from .kgm import province_id

FOLDER = RAW / "diyanet" if (RAW / "diyanet").exists() else Path("C:/veri-ham/diyanet")
CODE = re.compile(r"^TR[0-9ABC][0-9]{2}$")
FILES = {
    "personnel_series": "1_1_personel_sayisi_2023.xls",
    "personnel_province": (
        "1_3_istatistiki_bolge_birimleri_siniflamasina_ve_kadrolarina_gore_"
        "personel_sayisi_Muftulukler_2023.xlsx"
    ),
    "mosque_series": "2_1_cami_sayisi_2023.xls",
    "mosque_province": (
        "2_2_istatistiki_bolge_birimleri_siniflamasina_gore_cami_sayisi_2023.xls"
    ),
    "quran_province": (
        "3_3_istatistiki_bolge_birimleri_siniflamasina_gore_Kuran_kursu_kursiyer_ve_"
        "bitiren_kursiyer_sayisi_2023.xlsx"
    ),
    "converts": "4_1_cesitli_ozelliklerine_gore_ihtida_edenler_2023.xlsx",
    "pilgrims": "5_1_yas_gruplarına_göre hacca_ve_umreye_gidenlerin_sayisi_2023.xlsx",
    "budget": "6_1_2022-2023_yillari_butce_harcamaları_2023.xlsx",
}
#: personnel table 1.3: staff group -> the columns to add up (0-indexed)
STAFF_COLUMNS = {
    "total": (2,),
    "contract": (3,),
    "permanent": (4,),
    "general_admin": (6,),
    "mufti": (7,),
    "preacher": (15,),
    "quran_teacher": (18, 19),
    "imam": (20, 21),
    "muezzin": (22, 23),
}
#: Qur'an course table 3.3: (indicator, dims) -> column
QURAN_COLUMNS = {
    ("diyanet_quran_courses", ""): 2,
    ("diyanet_quran_attendance", ""): 3,
    ("diyanet_quran_hafiz", "sex=total"): 4,
    ("diyanet_quran_hafiz", "sex=male"): 5,
    ("diyanet_quran_hafiz", "sex=female"): 6,
    ("diyanet_quran_graduates", "sex=total"): 8,
    ("diyanet_quran_graduates", "sex=male"): 9,
    ("diyanet_quran_graduates", "sex=female"): 10,
    ("diyanet_quran_graduates_education", "education_level=illiterate"): 11,
    ("diyanet_quran_graduates_education", "education_level=literate_no_school"): 12,
    ("diyanet_quran_graduates_education", "education_level=preschool"): 13,
    ("diyanet_quran_graduates_education", "education_level=primary"): 14,
    ("diyanet_quran_graduates_education", "education_level=basic_8_year"): 15,
    ("diyanet_quran_graduates_education", "education_level=upper_secondary"): 16,
    ("diyanet_quran_graduates_education", "education_level=higher"): 17,
    ("diyanet_quran_graduates_age", "age_group=under_15"): 19,
    ("diyanet_quran_graduates_age", "age_group=15_17"): 20,
    ("diyanet_quran_graduates_age", "age_group=18_22"): 21,
    ("diyanet_quran_graduates_age", "age_group=23_44"): 22,
    ("diyanet_quran_graduates_age", "age_group=45_plus"): 23,
}
#: converts table 4.1: row label -> (indicator, dim value)
CONVERT_ROWS = {
    "Hristiyan": ("diyanet_converts", "former_belief=christian"),
    "Musevi": ("diyanet_converts", "former_belief=jewish"),
    "Budist": ("diyanet_converts", "former_belief=buddhist"),
    "Ateist": ("diyanet_converts", "former_belief=atheist"),
    "Diğer dinler": ("diyanet_converts", "former_belief=other"),
    "İnceleme ve araştırma": ("diyanet_converts_reason", "convert_reason=study"),
    "Evlenme": ("diyanet_converts_reason", "convert_reason=marriage"),
    "Tavsiye-Yönlendirme": ("diyanet_converts_reason", "convert_reason=advice"),
    "Etkilenme": ("diyanet_converts_reason", "convert_reason=influence"),
    "0-20": ("diyanet_converts_age", "age_group=0_20"),
    "21-30": ("diyanet_converts_age", "age_group=21_30"),
    "31-40": ("diyanet_converts_age", "age_group=31_40"),
    "41-50": ("diyanet_converts_age", "age_group=41_50"),
    "51-60": ("diyanet_converts_age", "age_group=51_60"),
    "61+": ("diyanet_converts_age", "age_group=61_plus"),
}
#: budget table 6.1: English row label -> expenditure type
BUDGET_ROWS = {
    "Personnel Expenditures": "personnel",
    "Other Current Expenditures": "other_current",
    "Investments": "investment",
    "Transfers": "transfer",
    "General Total": "total",
}
SEXES = ("total", "male", "female")
UNITS = {
    "diyanet_mosques": "facility",
    "diyanet_personnel": "person",
    "diyanet_quran_courses": "facility",
    "diyanet_quran_attendance": "person",
    "diyanet_quran_hafiz": "person",
    "diyanet_quran_graduates": "person",
    "diyanet_quran_graduates_education": "person",
    "diyanet_quran_graduates_age": "person",
    "diyanet_pilgrims": "person",
    "diyanet_converts": "person",
    "diyanet_converts_reason": "person",
    "diyanet_converts_age": "person",
    "diyanet_budget": "try",
}
Row = tuple[str, str, str, int]  # indicator, area, dims, year


def number(cell) -> float | None:
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    text = str(cell).strip().replace(" ", "").replace(",", ".")
    return float(text) if re.fullmatch(r"-?\d+(\.\d+)?", text) else None


def rows_of(name: str) -> list[list]:
    path = FOLDER / name
    if not path.exists():
        raise FileNotFoundError(f"{path} yok (scripts/fetch_diyanet.py)")
    if path.suffix == ".xlsx":
        import openpyxl

        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book[book.sheetnames[0]]
        # 6.1 is saved with the whole sheet allocated; the tables are small
        out = [
            list(row)
            for row in sheet.iter_rows(max_row=120, max_col=30, values_only=True)
        ]
        book.close()
        return out
    import xlrd

    sheet = xlrd.open_workbook(path).sheet_by_index(0)
    return [sheet.row_values(index) for index in range(sheet.nrows)]


def provincial(name: str, side_by_side: bool = False) -> dict[str, list]:
    """{area_id or "TR": row} of a table printed by level-3 statistical region."""
    out: dict[str, list] = {}
    for row in rows_of(name):
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        starts = (
            [0]
            if not side_by_side
            else [
                index
                for index, cell in enumerate(cells)
                if CODE.match(cell) or cell == "TR"
            ]
        )
        for start in starts:
            if start >= len(cells):
                continue
            code = cells[start]
            if code == "TR":
                out["TR"] = list(row[start:])
            elif CODE.match(code):
                out[province_id(cells[start + 1])] = list(row[start:])
    return out


def read_mosques() -> dict[Row, float]:
    out: dict[Row, float] = {}
    for row in rows_of(FILES["mosque_series"]):
        year, value = (
            number(row[0] if row else None),
            number(row[1] if len(row) > 1 else None),
        )
        if year and 2000 < year < 2100 and value:
            out[("diyanet_mosques", "TR", "", int(year))] = value
    table = provincial(FILES["mosque_province"], side_by_side=True)
    printed = number(table["TR"][2])
    for area, row in table.items():
        if area != "TR":
            out[("diyanet_mosques", area, "", 2023)] = number(row[2])
    check(
        "cami 2023",
        sum(
            v
            for (i, a, _d, _y), v in out.items()
            if i == "diyanet_mosques" and a != "TR"
        ),
        printed,
    )
    return out


def read_personnel() -> dict[Row, float]:
    out: dict[Row, float] = {}
    for row in rows_of(FILES["personnel_series"]):
        year = number(row[0] if row else None)
        if not year or not 2000 < year < 2100:
            continue
        total, permanent, contract = (number(row[i]) for i in (1, 2, 4))
        if total is None:
            continue
        for group, value in (
            ("total", total),
            ("permanent", permanent),
            ("contract", contract),
        ):
            if value is not None:
                out[("diyanet_personnel", "TR", f"staff_group={group}", int(year))] = (
                    value
                )
        check(f"personel {int(year)}", (permanent or 0) + (contract or 0), total)
    table = provincial(FILES["personnel_province"])
    for area, row in table.items():
        if area == "TR":
            continue
        for group, columns in STAFF_COLUMNS.items():
            values = [number(row[c]) for c in columns if c < len(row)]
            if any(v is not None for v in values):
                out[
                    ("diyanet_personnel_muftiships", area, f"staff_group={group}", 2023)
                ] = sum(v for v in values if v is not None)
    # the provinces' own "genel toplam" leaves out the general administration staff that the
    # Türkiye row of the same table counts: 127,194 + 9,190 = 136,384 in 2023
    group = lambda name: sum(
        v
        for (i, _a, d, _y), v in out.items()
        if i == "diyanet_personnel_muftiships" and d == f"staff_group={name}"
    )
    check(
        "müftülük personeli 2023",
        group("total") + group("general_admin"),
        number(table["TR"][2]),
    )
    return out


def read_quran() -> dict[Row, float]:
    out: dict[Row, float] = {}
    table = provincial(FILES["quran_province"])
    for area, row in table.items():
        if area == "TR":
            continue
        for (indicator, dims), column in QURAN_COLUMNS.items():
            value = number(row[column]) if column < len(row) else None
            if value is not None:
                out[(indicator, area, dims, 2023)] = value
    total = lambda indicator, dims: sum(
        v for (i, _a, d, _y), v in out.items() if i == indicator and d == dims
    )
    check("kurs sayısı", total("diyanet_quran_courses", ""), number(table["TR"][2]))
    check("kursiyer", total("diyanet_quran_attendance", ""), number(table["TR"][3]))
    graduates = total("diyanet_quran_graduates", "sex=total")
    check(
        "bitiren = erkek + kadın",
        total("diyanet_quran_graduates", "sex=male")
        + total("diyanet_quran_graduates", "sex=female"),
        graduates,
    )
    check(
        "bitiren = eğitim durumu",
        sum(
            v
            for (i, _a, _d, _y), v in out.items()
            if i == "diyanet_quran_graduates_education"
        ),
        graduates,
    )
    check(
        "bitiren = yaş grubu",
        sum(
            v
            for (i, _a, _d, _y), v in out.items()
            if i == "diyanet_quran_graduates_age"
        ),
        graduates,
    )
    return out


def read_pilgrims() -> dict[Row, float]:
    out: dict[Row, float] = {}
    printed: dict[str, float] = {}
    for row in rows_of(FILES["pilgrims"]):
        label = str(row[0]).strip() if row and row[0] else ""
        if not re.fullmatch(
            r"-?\d+(-\d+|\+)?|Toplam\nTotal|Toplam", label.replace(" ", "")
        ):
            continue
        band = (
            "total"
            if label.startswith("Toplam")
            else label.replace("-20", "under_20")
            .replace("+", "_plus")
            .replace("-", "_")
        )
        for kind, offset in (("hajj", 1), ("umrah", 5)):
            for index, sex in enumerate(SEXES):
                value = (
                    number(row[offset + index]) if offset + index < len(row) else None
                )
                if value is None:
                    continue
                if band == "total":
                    printed[f"{kind}_{sex}"] = value
                else:
                    out[
                        (
                            "diyanet_pilgrims",
                            "TR",
                            f"pilgrimage={kind};sex={sex};age_group={band}",
                            2023,
                        )
                    ] = value
    for key, value in printed.items():
        kind, sex = key.rsplit("_", 1)
        check(
            f"{kind} {sex}",
            sum(
                v
                for (_i, _a, d, _y), v in out.items()
                if f"pilgrimage={kind};sex={sex};" in d
            ),
            value,
        )
    return out


def read_converts() -> dict[Row, float]:
    out: dict[Row, float] = {}
    printed = None
    for row in rows_of(FILES["converts"]):
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        if (
            printed is None
            and number(cells[2] if len(cells) > 2 else None)
            and not cells[1]
        ):
            printed = number(cells[2])
        label = cells[1] if len(cells) > 1 else ""
        found = next(
            (value for key, value in CONVERT_ROWS.items() if label.startswith(key)),
            None,
        )
        if found is None:
            continue
        indicator, dim = found
        for index, sex in enumerate(SEXES):
            value = number(row[2 + index]) if 2 + index < len(row) else None
            if value is not None:
                out[(indicator, "TR", f"{dim};sex={sex}", 2023)] = value
    for indicator in (
        "diyanet_converts",
        "diyanet_converts_reason",
        "diyanet_converts_age",
    ):
        check(
            indicator,
            sum(
                v
                for (i, _a, d, _y), v in out.items()
                if i == indicator and d.endswith("sex=total")
            ),
            printed,
        )
    return out


def read_budget() -> dict[Row, float]:
    out: dict[Row, float] = {}
    for row in rows_of(FILES["budget"]):
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        label = next((cell for cell in cells[:2] if cell), "")
        kind = BUDGET_ROWS.get(label)
        if kind is None:
            continue
        for index, year in enumerate((2022, 2023)):
            value = number(row[2 + index]) if 2 + index < len(row) else None
            if value is not None:
                out[("diyanet_budget", "TR", f"expenditure_type={kind}", year)] = value
    for year in (2022, 2023):
        parts = sum(
            v
            for (_i, _a, d, y), v in out.items()
            if y == year and d != "expenditure_type=total"
        )
        check(
            f"bütçe {year}",
            parts,
            out[("diyanet_budget", "TR", "expenditure_type=total", year)],
        )
    return out


def check(what: str, read: float | None, printed: float | None) -> None:
    if printed is None or read is None:
        raise ValueError(f"Diyanet {what}: toplam okunamadı")
    if abs(read - printed) > max(1.0, printed * 0.001):
        raise ValueError(f"Diyanet {what}: okunan {read:,.0f}, basılı {printed:,.0f}")


def load_all() -> dict[Row, float]:
    out: dict[Row, float] = {}
    for reader in (
        read_mosques,
        read_personnel,
        read_quran,
        read_pilgrims,
        read_converts,
        read_budget,
    ):
        out.update(reader())
    return out


_CACHE: dict[Row, float] = {}


class Diyanet:
    source_id = "diyanet"
    indicator_id = ""

    def fetch(self) -> Path:
        return FOLDER

    def parse(self, raw: Path) -> pl.DataFrame:
        if not _CACHE:
            _CACHE.update(load_all())
        # the muftiship personnel are the same indicator at province level
        wanted = {self.indicator_id}
        if self.indicator_id == "diyanet_personnel":
            wanted.add("diyanet_personnel_muftiships")
        records = [
            {
                "area_id": area,
                "area_level": "country" if area == "TR" else "province",
                "period_start": dt.date(year, 1, 1),
                "dims": dims,
                "value": value,
            }
            for (indicator, area, dims, year), value in _CACHE.items()
            if indicator in wanted
        ]
        return pl.DataFrame(
            records, schema_overrides={"value": pl.Float64}
        ).with_columns(
            pl.lit(self.indicator_id).alias("indicator_id"),
            pl.lit("annual").alias("frequency"),
            pl.lit(UNITS[self.indicator_id]).alias("unit"),
            pl.lit("measured").alias("quality_flag"),
            pl.lit("2026-09").alias("vintage"),
            pl.lit("diyanet").alias("source_id"),
            pl.lit(dt.date(2026, 9, 16)).alias("retrieved_at"),
        )


DIYANET_ADAPTERS = {
    indicator: type(f"Diyanet_{indicator}", (Diyanet,), {"indicator_id": indicator})
    for indicator in UNITS
}
