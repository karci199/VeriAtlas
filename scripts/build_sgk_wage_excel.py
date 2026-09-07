"""İl bazında ortalama günlük kazanç (4/1-a) — 2025 seviyesi, kamu/özel ve cinsiyet farkı 2014 vs 2025.

SGK İstatistik Yıllığı, Tablo 1.17: "4/1-a Kapsamındaki İşyeri, Zorunlu Sigortalılar ile
Prime Esas Ortalama Günlük Kazançların İl, Sektör ve Cinsiyete Göre Dağılımı" — il başına
kamu/özel ve erkek/kadın ortalama günlük kazanç (TL), 2014 ve 2025'te aynı sütun düzeni.

**Nominal seviye karşılaştırılmıyor.** 2014→2025 arası TL bu denli değer kaybetmişken
("60 TL" ile "1500 TL" yan yana anlamsız — deflatör yok, elimizde yıllık TÜFE serisi bu
depoda değil) iki yılın TL tutarını yan yana koymak yanıltıcı olurdu. Onun yerine yalnız
**aynı yıl içindeki oranlar** taşınıyor — kamu/özel kazanç oranı ve kadın/erkek kazanç
oranı — çünkü bunlar enflasyondan bağımsız, iki yıl arasında karşılaştırılabilir.

Run:  uv run python scripts/build_sgk_wage_excel.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import polars as pl
import xlsxwriter

sys.path.insert(0, "src")
sys.path.insert(0, "scripts")

from build_analysis_books import notes_sheet, ranked, styles

from veriatlas.config import PUBLIC

SCRATCH = Path(
    r"C:\Users\katan\AppData\Local\Temp\claude\C--veri--claude-worktrees-nerede-kaldik-4998d7"
    r"\289f51e5-ac10-4275-b09d-211258ba36af\scratchpad"
)
TARGET = PUBLIC.parent / "cikti" / "analiz-sgk-gunluk-kazanc-il-2025.xlsx"


def read_year(path: Path) -> pl.DataFrame:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["TABLO-1.17 "] if "TABLO-1.17 " in wb.sheetnames else wb["TABLO-1.17"]
    rows = []
    for row in ws.iter_rows(min_row=7, max_row=95, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        sigortali = row[13]
        kamu, ozel, erkek, kadin, toplam = row[16], row[17], row[18], row[19], row[20]
        if toplam is None:
            continue
        rows.append(
            {
                "il": str(il).strip(),
                "plaka": plate,
                "sigortali": sigortali,
                "kazanc_toplam": toplam,
                "kazanc_kamu": kamu,
                "kazanc_ozel": ozel,
                "kazanc_erkek": erkek,
                "kazanc_kadin": kadin,
            }
        )
    return pl.DataFrame(rows)


def main() -> None:
    d2025 = read_year(next((SCRATCH / "sgk2025").glob("*.xlsx")))
    d2014 = read_year(
        SCRATCH / "sgk2014" / "sgk_2014"
        / "2014 YILLIK BÖLÜM 1 İsyeri ve Sigortalılara Ait İstatistikler.xlsx"
    )
    for name, d in (("2025", d2025), ("2014", d2014)):
        if len(d) != 81:
            print(f"uyari [{name}]: 81 il beklenirken {len(d)} satir okundu")

    # 2025 seviyesi — kamu/özel oranı, kadın/erkek oranı da eklendi.
    level = ranked(
        d2025.with_columns(
            (pl.col("kazanc_kamu") / pl.col("kazanc_ozel")).alias("kamu_ozel_orani"),
            (pl.col("kazanc_kadin") / pl.col("kazanc_erkek")).alias("kadin_erkek_orani"),
        ).select(
            "il", "plaka", "kazanc_toplam", "kazanc_kamu", "kazanc_ozel",
            "kamu_ozel_orani", "kazanc_erkek", "kazanc_kadin", "kadin_erkek_orani",
        ),
        "kazanc_toplam",
    )

    # Yalnız oranlar, iki yıl — enflasyondan bağımsız karşılaştırma.
    def ratios(d: pl.DataFrame, tag: str) -> pl.DataFrame:
        return d.select(
            "il", "plaka",
            (pl.col("kazanc_kamu") / pl.col("kazanc_ozel")).alias(f"kamu_ozel_{tag}"),
            (pl.col("kazanc_kadin") / pl.col("kazanc_erkek")).alias(f"kadin_erkek_{tag}"),
        )

    compare = ratios(d2014, "2014").join(ratios(d2025, "2025").drop("plaka"), on="il").with_columns(
        (pl.col("kamu_ozel_2025") - pl.col("kamu_ozel_2014")).alias("kamu_ozel_fark"),
        (pl.col("kadin_erkek_2025") - pl.col("kadin_erkek_2014")).alias("kadin_erkek_fark"),
    )
    compare = ranked(compare, "kadin_erkek_fark")

    # region Yazma
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(str(TARGET))
    style = styles(book)
    left = book.add_format({"align": "left", "valign": "vcenter", "indent": 1})
    tl = book.add_format({"num_format": "#,##0", "align": "center", "valign": "vcenter"})
    ratio_fmt = book.add_format({"num_format": "0.00", "align": "center", "valign": "vcenter"})

    headers = {
        "il": "İl", "plaka": "Plaka", "sira": "Sıra",
        "sigortali": "Zorunlu sigortalı",
        "kazanc_toplam": "Ortalama günlük kazanç (TL) · 2025",
        "kazanc_kamu": "Kamu · 2025", "kazanc_ozel": "Özel · 2025",
        "kazanc_erkek": "Erkek · 2025", "kazanc_kadin": "Kadın · 2025",
        "kamu_ozel_orani": "Kamu/Özel oranı · 2025",
        "kadin_erkek_orani": "Kadın/Erkek oranı · 2025",
        "kamu_ozel_2014": "Kamu/Özel oranı · 2014", "kamu_ozel_2025": "Kamu/Özel oranı · 2025",
        "kamu_ozel_fark": "Kamu/Özel oranı farkı",
        "kadin_erkek_2014": "Kadın/Erkek oranı · 2014", "kadin_erkek_2025": "Kadın/Erkek oranı · 2025",
        "kadin_erkek_fark": "Kadın/Erkek oranı farkı",
    }
    tl_cols = ("kazanc_toplam", "kazanc_kamu", "kazanc_ozel", "kazanc_erkek", "kazanc_kadin")
    ratio_cols = tuple(c for c in compare.columns if c not in ("il", "plaka", "sira")) + (
        "kamu_ozel_orani", "kadin_erkek_orani",
    )
    formats = {
        "il": left, "plaka": style["text"], "sira": style["text"],
        "sigortali": style["count"],
        **{c: tl for c in tl_cols},
        **{c: ratio_fmt for c in ratio_cols},
    }
    widths = {"il": 22, "plaka": 8, "sira": 7}

    def write_sheet(name: str, data: pl.DataFrame) -> None:
        page = book.add_worksheet(name)
        page.freeze_panes(1, 1)
        page.set_row(0, 40)
        columns = data.columns
        for index, column in enumerate(columns):
            page.set_column(index, index, widths.get(column, 18), formats.get(column, style["text"]))
        rows_ = [list(r) for r in data.rows()]
        page.add_table(
            0, 0, max(len(rows_), 1), len(columns) - 1,
            {
                "data": rows_,
                "columns": [
                    {"header": headers.get(c, c), "header_format": style["head"], "format": formats.get(c, style["text"])}
                    for c in columns
                ],
                "banded_rows": True,
            },
        )

    write_sheet("2025 seviyesi", level)
    write_sheet("Kamu-Özel ve Cinsiyet farkı", compare)

    notes_sheet(
        book,
        [
            "İl bazında ortalama günlük kazanç (4/1-a) — 2025 seviyesi, kamu/özel ve",
            "cinsiyet farkı 2014 vs 2025",
            "",
            "Yalnız 4/1-a (işçi statüsü) kapsıyor. Kazanç = prime esas ortalama günlük",
            "kazanç (TL), SGK'nın kendi hesapladığı rakam.",
            "",
            "NOMİNAL KARŞILAŞTIRMA YOK",
            "2014'te TL bugünkünün çok altında bir değere sahipti (~60 TL/gün ortalama",
            "kazanç, 2025'te ~1500 TL) — deflatör (TÜFE serisi) bu depoda yok, o yüzden",
            "iki yılın TL tutarını yan yana koymak yanıltıcı olurdu. Bunun yerine yalnız",
            "aynı yıl içindeki ORANLAR taşındı — kamu/özel ve kadın/erkek kazanç oranı —",
            "çünkü bunlar enflasyondan bağımsız, iki yıl arasında karşılaştırılabilir.",
            "",
            "SAYFALAR",
            "· 2025 seviyesi — il başına TL tutarları, yalnız 2025, ortalama kazanca",
            "  göre sıralı.",
            "· Kamu-Özel ve Cinsiyet farkı — yalnız oranlar, iki yıl, fark. Kadın/Erkek",
            "  oranı farkına göre sıralı (pozitif = kadın lehine daralma).",
            "",
            "KAYNAK",
            "SGK İstatistik Yıllığı, Tablo 1.17 (2014 ve 2025 sürümleri,",
            "sgk.gov.tr/Istatistik/Yillik).",
        ],
    )
    book.close()
    # endregion

    print("yazildi:", TARGET)


if __name__ == "__main__":
    main()
