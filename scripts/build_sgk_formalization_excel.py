"""İllere göre zorunlu sigortalılık oranı (20-64 nüfusa göre), cinsiyet kırılımlı — 2014 vs 2025.

Sigortalı sayıp bırakmak yerine bir paydaya oturtuyor: kaç kişi çalışıyor değil,
**çalışma çağındaki nüfusun ne kadarı kayıtlı sigortalı**. Pay = 4/1-a + 4/1-b + 4/1-c
zorunlu sigortalı toplamı (SGK İstatistik Yıllığı, Tablo 1.8 — üç statüyü tek tabloda,
il ve cinsiyete göre veren tek tablo). Payda = 20-64 yaş nüfusu (population.csv.gz),
cinsiyete göre.

**20-64 tercih edildi, 25-60 değil:** iki nedenle. Birincisi veri — nüfus dosyası 5'lik
yaş grubunda (20-24, ..., 60-64), 20-64 tam grup sınırlarına oturuyor, 25-60 bir grubu
ortadan kesip tek yaş dosyasına ihtiyaç duyardı. İkincisi kavramsal — 25-60 erken
başlayan mavi yakalı/çırak işçiliği (20-24 yaş) ve 60-64 arası hâlâ aktif çalışanları
dışarıda bırakıp paydayı gereğinden çok küçültüyor, oranı yapay şişiriyor.

**Sınır:** pay yalnız "zorunlu" sigortalı (isteğe bağlı, stajyer, çırak, yurtdışı hariç)
— Tablo 1.8 bu kırılımı yalnız zorunlu için veriyor. Toplam aktif sigortalıdan (Tablo
1.7) birkaç puan düşük çıkar, cinsiyet kırılımının bedeli bu.

Run:  uv run python scripts/build_sgk_formalization_excel.py
"""

from __future__ import annotations

import csv
import gzip
import sys
from pathlib import Path

import openpyxl
import polars as pl
import xlsxwriter

sys.path.insert(0, "src")
sys.path.insert(0, "scripts")

from build_analysis_books import notes_sheet, ranked, styles

from veriatlas.config import PUBLIC

SCRATCH = Path(
    r"C:\Users\katan\AppData\Local\Temp\claude\C--veri--claude-worktrees-nerede-kaldik-4998d7"
    r"\289f51e5-ac10-4275-b09d-211258ba36af\scratchpad"
)
TARGET = PUBLIC.parent / "cikti" / "analiz-sgk-formalizasyon-20-64-il-yillara-gore.xlsx"

WORKING_AGES = {"20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "50-54", "55-59", "60-64"}


def read_2014() -> pl.DataFrame:
    path = SCRATCH / "sgk2014" / "sgk_2014" / "2014 YILLIK BÖLÜM 1 İsyeri ve Sigortalılara Ait İstatistikler.xlsx"
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["Tablo-1.8"]
    rows = []
    for row in ws.iter_rows(min_row=6, max_row=90, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        a_k, a_e = row[2], row[3]
        b1_k, b1_e, b1_t = row[5], row[6], row[7]
        b2_k, b2_e, b2_t = row[8], row[9], row[10]
        c_k, c_e = row[11], row[12]
        if a_k is None:
            continue
        rows.append(
            {
                "il": str(il).strip(),
                "plaka": plate,
                "sigortali_e": a_e + b1_e + b2_e + c_e,
                "sigortali_k": a_k + b1_k + b2_k + c_k,
            }
        )
    return pl.DataFrame(rows).with_columns(
        (pl.col("sigortali_e") + pl.col("sigortali_k")).alias("sigortali_t")
    )


def read_2025() -> pl.DataFrame:
    path = next((SCRATCH / "sgk2025").glob("*.xlsx"))
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["Tablo-1.8"]
    rows = []
    for row in ws.iter_rows(min_row=8, max_row=90, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        a_e, a_k = row[2], row[3]
        b_e, b_k = row[14], row[15]
        c_e, c_k = row[17], row[18]
        if a_e is None:
            continue
        rows.append(
            {
                "il": str(il).strip(),
                "plaka": plate,
                "sigortali_e": a_e + b_e + c_e,
                "sigortali_k": a_k + b_k + c_k,
            }
        )
    return pl.DataFrame(rows).with_columns(
        (pl.col("sigortali_e") + pl.col("sigortali_k")).alias("sigortali_t")
    )


def population_20_64(year: int) -> dict[int, dict[str, float]]:
    out: dict[int, dict[str, float]] = {}
    with gzip.open(PUBLIC / "population.csv.gz", "rt", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["level"] != "province" or int(row["year"]) != year:
                continue
            if row["age"] not in WORKING_AGES:
                continue
            plate = int(row["area_id"].split("-")[1])
            bucket = out.setdefault(plate, {"male": 0.0, "female": 0.0})
            bucket[row["sex"]] += float(row["value"])
    return out


def main() -> None:
    sig14 = read_2014().rename({"il": "il14"})
    sig25 = read_2025()
    if len(sig14) != 81 or len(sig25) != 81:
        print(f"uyari: 81 il beklenirken 2014={len(sig14)} 2025={len(sig25)}")

    pop14 = population_20_64(2014)
    pop25 = population_20_64(2025)

    rows = []
    for plate in sig25["plaka"].to_list():
        il = sig25.filter(pl.col("plaka") == plate)["il"][0]
        s14 = sig14.filter(pl.col("plaka") == plate)
        s25 = sig25.filter(pl.col("plaka") == plate)
        p14 = pop14.get(plate, {"male": 0.0, "female": 0.0})
        p25 = pop25.get(plate, {"male": 0.0, "female": 0.0})

        e14, k14, t14 = s14["sigortali_e"][0], s14["sigortali_k"][0], s14["sigortali_t"][0]
        e25, k25, t25 = s25["sigortali_e"][0], s25["sigortali_k"][0], s25["sigortali_t"][0]
        pop_e14, pop_k14 = p14["male"], p14["female"]
        pop_e25, pop_k25 = p25["male"], p25["female"]

        rows.append(
            {
                "il": il,
                "plaka": plate,
                "oran_t_2014": t14 / (pop_e14 + pop_k14) if (pop_e14 + pop_k14) else None,
                "oran_t_2025": t25 / (pop_e25 + pop_k25) if (pop_e25 + pop_k25) else None,
                "oran_e_2014": e14 / pop_e14 if pop_e14 else None,
                "oran_e_2025": e25 / pop_e25 if pop_e25 else None,
                "oran_k_2014": k14 / pop_k14 if pop_k14 else None,
                "oran_k_2025": k25 / pop_k25 if pop_k25 else None,
                "sigortali_t_2014": t14,
                "sigortali_t_2025": t25,
                "nufus20_64_2014": pop_e14 + pop_k14,
                "nufus20_64_2025": pop_e25 + pop_k25,
            }
        )

    frame = pl.DataFrame(rows).with_columns(
        (pl.col("oran_t_2025") - pl.col("oran_t_2014")).alias("oran_t_fark"),
        # Cinsiyet farkı — erkek oranı kadından ne kadar yüksek, o yılda.
        (pl.col("oran_e_2014") - pl.col("oran_k_2014")).alias("cinsiyet_farki_2014"),
        (pl.col("oran_e_2025") - pl.col("oran_k_2025")).alias("cinsiyet_farki_2025"),
    ).with_columns(
        # Fark ne kadar daraldı — pozitif = kadın lehine iyileşme.
        (pl.col("cinsiyet_farki_2014") - pl.col("cinsiyet_farki_2025")).alias("iyilesme")
    )

    # Sade sayfa: yalnız toplam sigortalılık oranı, iki yıl ve fark.
    toplam = ranked(
        frame.select("il", "plaka", "oran_t_2014", "oran_t_2025", "oran_t_fark"),
        "oran_t_fark",
    )

    # Sade sayfa: cinsiyet farkının kendisi ve ne kadar daraldığı.
    cinsiyet = ranked(
        frame.select(
            "il", "plaka",
            "cinsiyet_farki_2014", "cinsiyet_farki_2025", "iyilesme",
        ),
        "iyilesme",
    )

    tr = pl.DataFrame(
        [
            {
                "yil": 2014,
                "sigortali_t": frame["sigortali_t_2014"].sum(),
                "nufus20_64": frame["nufus20_64_2014"].sum(),
                "oran_t": frame["sigortali_t_2014"].sum() / frame["nufus20_64_2014"].sum(),
            },
            {
                "yil": 2025,
                "sigortali_t": frame["sigortali_t_2025"].sum(),
                "nufus20_64": frame["nufus20_64_2025"].sum(),
                "oran_t": frame["sigortali_t_2025"].sum() / frame["nufus20_64_2025"].sum(),
            },
        ]
    )

    # region Yazma
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(str(TARGET))
    style = styles(book)
    left = book.add_format({"align": "left", "valign": "vcenter", "indent": 1})

    headers = {
        "il": "İl", "plaka": "Plaka", "sira": "Sıra", "yil": "Yıl",
        "sigortali_t": "Zorunlu sigortalı (toplam)", "nufus20_64": "Nüfus 20-64",
        "oran_t": "Sigortalılık oranı",
        "oran_t_2014": "Oran · 2014", "oran_t_2025": "Oran · 2025",
        "oran_t_fark": "Fark",
        "cinsiyet_farki_2014": "Cinsiyet farkı (E−K) · 2014",
        "cinsiyet_farki_2025": "Cinsiyet farkı (E−K) · 2025",
        "iyilesme": "Fark daralması (kadın lehine)",
    }
    percent_cols = {
        "oran_t_2014", "oran_t_2025", "oran_t_fark",
        "cinsiyet_farki_2014", "cinsiyet_farki_2025", "iyilesme", "oran_t",
    }
    formats = {
        "il": left, "plaka": style["text"], "sira": style["text"], "yil": style["text"],
        **{c: style["percent"] for c in percent_cols},
        "sigortali_t": style["count"], "nufus20_64": style["count"],
    }
    widths = {"il": 22, "plaka": 8, "sira": 7, "yil": 8}

    def write_sheet(name: str, data: pl.DataFrame) -> None:
        page = book.add_worksheet(name)
        page.freeze_panes(1, 1)
        page.set_row(0, 40)
        columns = data.columns
        for index, column in enumerate(columns):
            page.set_column(index, index, widths.get(column, 18), formats.get(column, style["text"]))
        rows_ = [list(r) for r in data.rows()]
        page.add_table(
            0, 0, max(len(rows_), 1), len(columns) - 1,
            {
                "data": rows_,
                "columns": [
                    {"header": headers.get(c, c), "header_format": style["head"], "format": formats.get(c, style["text"])}
                    for c in columns
                ],
                "banded_rows": True,
            },
        )

    write_sheet("Toplam oran", toplam)
    write_sheet("Cinsiyet farkı", cinsiyet)
    write_sheet("Türkiye", tr)

    notes_sheet(
        book,
        [
            "İllere göre zorunlu sigortalılık oranı (20-64 nüfusa göre) — 2014 vs 2025",
            "",
            "Oran = zorunlu sigortalı (4/1-a + 4/1-b + 4/1-c) / 20-64 yaş nüfusu, cinsiyete",
            "göre ayrı hesaplandı. Kayıtlı çalışan oranına yakın bir ölçü ama TÜİK'in resmi",
            "işgücüne katılım oranı DEĞİL — emekliler, işsizler, ev hanımları, öğrenciler",
            "paydada; sigortalı olmayan bakmakla yükümlü nüfus da.",
            "",
            "20-64 seçildi, 25-60 değil: nüfus dosyası 5'lik yaş grubunda (20-24...60-64),",
            "20-64 tam grup sınırına oturuyor. 25-60 bir grubu ortadan keser, ayrıca 20-24",
            "yaşındaki çırak/genç işçiyi ve 60-64 arası hâlâ aktif çalışanı dışarıda bırakıp",
            "paydayı yapay küçültür.",
            "",
            "Pay yalnız 'zorunlu' sigortalı — isteğe bağlı, stajyer, çırak, yurtdışı hariç,",
            "çünkü cinsiyet kırılımı SGK Tablo 1.8'de yalnız bu kırılım için veriliyor.",
            "Toplam aktif sigortalıdan (önceki dosyadaki Tablo 1.7 rakamları) birkaç puan",
            "düşük çıkar.",
            "",
            "SAYFALAR",
            "· Toplam oran — cinsiyet ayrımı olmadan, il başına tek oran. Fark sütununa",
            "  göre sıralı.",
            "· Cinsiyet farkı — erkek oranı kadın oranından kaç puan yüksek, iki yılda da.",
            "  'Fark daralması' = 2014'teki fark eksi 2025'teki fark: pozitifse kadın",
            "  lehine daralmış demektir. Sıralama bu sütuna göre — en çok daralan en üstte.",
            "",
            "KAYNAK",
            "SGK İstatistik Yıllığı, Tablo 1.8 (2014 ve 2025 sürümleri, sgk.gov.tr).",
            "Nüfus: population.csv.gz (TÜİK MEDAS, ADNKS), 20-64 yaş, 5'lik grup toplamı.",
        ],
    )
    book.close()
    # endregion

    print("yazildi:", TARGET)


if __name__ == "__main__":
    main()
