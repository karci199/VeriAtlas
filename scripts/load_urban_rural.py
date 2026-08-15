"""TÜİK's urban/rural class for every neighbourhood and village, into the registries.

The class is an *attribute of a place*, not a measurement of it, so it goes in the area
registry beside `bucak` and `municipality` rather than into the fact table. Storing it as
an indicator would mean writing the population a second time under a different name, and
two copies of one number is one copy too many: the split is computed by joining this
column to the population we already have.

Source: TÜİK's "Favori Raporlar" workbook, sheet `KENT-KIR SINIFLAMASI` — the only place
the classification is published per settlement. It is not in the MEDAS measure tree (all
39 measures were walked; it is not among them) and the SDMX flow `DF_ADNKS_T37` carries
the same classification only summed to provinces.

Three classes, kept as TÜİK's own three: `yogun_kent`, `orta_kent`, `kir`. Folding the
middle one into either side would be our judgement silently replacing the source's, and
it is 15,8% of the country — too big to hide.

Matched on the MEDAS record number, never on the name. The workbook's `MAHALLE KAYIT NO`
is the same number our registry stores as `medas_code`; matching Türkiye's neighbourhoods
by name would collapse the hundreds called "Merkez" and "Yeni" into each other.

One caveat travels with the column: the classification is a snapshot of 31 December 2025.
It has no history, so a place that urbanised in 2018 is labelled by what it is now, for
every year of the series.

Run:  uv run python scripts/load_urban_rural.py <FavoriRaporlar.xlsx>
"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import polars as pl

sys.path.insert(0, "src")

from veriatlas.config import RAW

DATA = Path("src/veriatlas/data")
SHEET = "KENT-KIR SINIFLAMASI"
SOURCE_ID = "tuik_favori"

#: TÜİK's labels to the codes we store. Turkish stays out of the data (K1).
CLASSES = {
    "YOĞUN KENT": "yogun_kent",
    "ORTA YOĞUN KENT": "orta_kent",
    "KIR": "kir",
}

# Column positions in the sheet. The header is split over two rows and the file carries
# no stable machine-readable names, so the positions are written down here rather than
# guessed at each run — a guess that silently shifts one column would relabel the country.
IL_KODU, ILCE_NO, KOY_NO, MAHALLE_NO = 1, 2, 3, 4
IL_ADI, ILCE_ADI, KOY_ADI = 5, 6, 9
TUR, SINIF = 10, 12


def read_sheet(path: Path) -> pl.DataFrame:
    """The classification sheet as rows of (kind, medas code, class)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise ValueError(f"{path.name} icinde '{SHEET}' sayfasi yok")

    rows = []
    for row in book[SHEET].iter_rows(values_only=True):
        label = row[SINIF]
        if not isinstance(row[0], int) or label not in CLASSES:
            continue  # title, explanation and header rows
        kind = row[TUR]
        code = row[KOY_NO] if kind == "KÖY" else row[MAHALLE_NO]
        if not isinstance(code, int):
            continue
        rows.append(
            {
                "level": "village" if kind == "KÖY" else "neighbourhood",
                "medas_code": code,
                "il": row[IL_ADI],
                "ilce": row[ILCE_ADI],
                "urban_rural": CLASSES[label],
            }
        )
    if not rows:
        raise ValueError("siniflama satiri okunamadi — sutun yerleri degismis olabilir")
    return pl.DataFrame(rows)


def attach(registry: Path, classes: pl.DataFrame, level: str) -> tuple[int, int]:
    """Write the class into one registry, keyed by MEDAS code. Returns (matched, total).

    A code that appears twice in the source would attach the wrong class to one of the
    two places, so the pairing is checked for uniqueness before it is trusted rather
    than after someone notices a strange map.
    """
    wanted = classes.filter(pl.col("level") == level).select(
        "medas_code", "urban_rural"
    )
    clashing = (
        wanted.group_by("medas_code")
        .agg(pl.col("urban_rural").n_unique().alias("n"))
        .filter(pl.col("n") > 1)
    )
    if len(clashing):
        raise ValueError(
            f"{level}: {len(clashing)} MEDAS kodu iki farkli sinifla geliyor, "
            "kod tek basina kimlik degil"
        )
    wanted = wanted.unique(subset="medas_code")

    frame = pl.read_csv(registry)
    if "urban_rural" in frame.columns:
        frame = frame.drop("urban_rural")
    joined = frame.join(wanted, on="medas_code", how="left")
    joined.write_csv(registry)
    return int(joined["urban_rural"].is_not_null().sum()), len(joined)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("kullanim: load_urban_rural.py <FavoriRaporlar.xlsx>")
    source = Path(sys.argv[1])
    if not source.exists():
        raise SystemExit(f"dosya yok: {source}")

    # The raw bytes are kept before anything is read out of them (K8): the parse can be
    # corrected and replayed later, the download cannot be taken again once TÜİK revises.
    RAW.mkdir(parents=True, exist_ok=True)
    kept = RAW / f"tuik-favori-raporlar-{datetime.now(UTC):%Y-%m-%d}.xlsx"
    if kept.resolve() != source.resolve():
        shutil.copy2(source, kept)
    digest = hashlib.sha256(kept.read_bytes()).hexdigest()[:16]

    classes = read_sheet(kept)
    counts = classes.group_by("level", "urban_rural").len().sort("level", "urban_rural")
    print(counts)

    report = {}
    for level, registry in (
        ("neighbourhood", DATA / "areas_tr_neighbourhoods.csv"),
        ("village", DATA / "areas_tr_villages.csv"),
    ):
        matched, total = attach(registry, classes, level)
        report[level] = (matched, total)
        print(f"{level}: {matched}/{total} eslesti ({100 * matched / total:.1f}%)")

    manifest = {
        "source_id": SOURCE_ID,
        "dataset": "urban_rural_classification",
        "file": kept.name,
        "sha256_16": digest,
        "rows": len(classes),
        "matched": {k: v[0] for k, v in report.items()},
        "of": {k: v[1] for k, v in report.items()},
        "as_of": "2025-12-31",
        "retrieved_at": datetime.now(UTC).strftime("%Y-%m-%d"),
    }
    with (RAW / "manifests.jsonl").open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(manifest, ensure_ascii=False) + "\n")
    print("manifest yazildi:", manifest["sha256_16"])


if __name__ == "__main__":
    main()
