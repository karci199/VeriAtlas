"""İl bazında emekli / toplam nüfus oranı — 2025.

Emekli sayısı (build_sgk_emekli_excel.py ile aynı tanım: malullük + yaşlılık aylığı,
4/1-a+4/1-b+4/1-c, ölüm aylığı hariç) / ilin toplam nüfusu. Yaş bandı zorlanmadı —
EYT (Emeklilikte Yaşa Takılanlar) yüzünden Türkiye'de emeklilik yaşı çok geniş bir
banda yayılmış durumda, 65+ gibi bir yaş grubuna bölmek emeklilerin büyük bölümünü
dışarıda bırakıp oranı yapay şişirirdi. Toplam nüfus paydası bu sorunu atlıyor: yaş
eşleştirmesi iddia etmiyor, yalnız "nüfusun ne kadarı emekli maaşı alıyor" sorusuna
cevap veriyor.

Run:  uv run python scripts/build_sgk_emekli_nufus_excel.py
"""

from __future__ import annotations

import csv
import gzip
import sys
from pathlib import Path

import openpyxl
import polars as pl
import xlsxwriter

sys.path.insert(0, "src")
sys.path.insert(0, "scripts")

from build_analysis_books import notes_sheet, ranked, styles

from veriatlas.config import PUBLIC

SCRATCH = Path(
    r"C:\Users\katan\AppData\Local\Temp\claude\C--veri--claude-worktrees-nerede-kaldik-4998d7"
    r"\289f51e5-ac10-4275-b09d-211258ba36af\scratchpad"
)
SGK2025 = SCRATCH / "sgk2025"
TARGET = PUBLIC.parent / "cikti" / "analiz-sgk-emekli-nufus-il-2025.xlsx"


def read_pensioners() -> dict[int, dict[str, float]]:
    path = SGK2025 / "YILLIK BÖLÜM 2 Aylık ve Gelir Alan İstatistikleri _2025.xlsx"
    wb = openpyxl.load_workbook(str(path), data_only=True)

    out: dict[int, dict[str, float]] = {}

    ws_a = wb["TABLO- 2.5"]
    for row in ws_a.iter_rows(min_row=8, max_row=95, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        emekli = (row[4] or 0) + (row[7] or 0)
        out.setdefault(plate, {}).update({"il": str(il).strip(), "emekli_a": emekli})

    ws_b = wb["TABLO- 2.19"]
    for row in ws_b.iter_rows(min_row=9, max_row=95, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        emekli = (row[4] or 0) + (row[7] or 0) + (row[18] or 0) + (row[21] or 0)
        out.setdefault(plate, {}).update({"emekli_b": emekli})

    ws_c = wb["TABLO-2.32"]
    for row in ws_c.iter_rows(min_row=8, max_row=95, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        emekli = (row[4] or 0) + (row[7] or 0) + (row[10] or 0)
        out.setdefault(plate, {}).update({"emekli_c": emekli})

    return out


def read_population(year: int) -> dict[int, float]:
    out: dict[int, float] = {}
    with gzip.open(PUBLIC / "population.csv.gz", "rt", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["level"] != "province" or int(row["year"]) != year:
                continue
            plate = int(row["area_id"].split("-")[1])
            out[plate] = out.get(plate, 0.0) + float(row["value"])
    return out


def main() -> None:
    pensioners = read_pensioners()
    population = read_population(2025)

    rows = []
    for plate, data in pensioners.items():
        if "il" not in data:
            continue
        emekli = data.get("emekli_a", 0) + data.get("emekli_b", 0) + data.get("emekli_c", 0)
        nufus = population.get(plate, 0)
        rows.append(
            {
                "il": data["il"],
                "plaka": plate,
                "emekli": emekli,
                "nufus": nufus,
                "emekli_nufus": emekli / nufus if nufus else None,
            }
        )
    frame = pl.DataFrame(rows)
    if len(frame) != 81:
        print(f"uyari: 81 il beklenirken {len(frame)} satir okundu")

    ranked_frame = ranked(
        frame.select("il", "plaka", "emekli", "nufus", "emekli_nufus"), "emekli_nufus"
    )

    tr_emekli = frame["emekli"].sum()
    tr_nufus = frame["nufus"].sum()
    tr = pl.DataFrame([{"emekli": tr_emekli, "nufus": tr_nufus, "emekli_nufus": tr_emekli / tr_nufus}])

    # region Yazma
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(str(TARGET))
    style = styles(book)
    left = book.add_format({"align": "left", "valign": "vcenter", "indent": 1})

    headers = {
        "il": "İl", "plaka": "Plaka", "sira": "Sıra",
        "emekli": "Emekli sayısı (malullük+yaşlılık)", "nufus": "Toplam nüfus",
        "emekli_nufus": "Emekli / Nüfus",
    }
    formats = {
        "il": left, "plaka": style["text"], "sira": style["text"],
        "emekli": style["count"], "nufus": style["count"],
        "emekli_nufus": style["percent"],
    }
    widths = {"il": 22, "plaka": 8, "sira": 7, "emekli": 20, "nufus": 14}

    def write_sheet(name: str, data: pl.DataFrame) -> None:
        page = book.add_worksheet(name)
        page.freeze_panes(1, 1)
        page.set_row(0, 40)
        columns = data.columns
        for index, column in enumerate(columns):
            page.set_column(index, index, widths.get(column, 16), formats.get(column, style["text"]))
        rows_ = [list(r) for r in data.rows()]
        page.add_table(
            0, 0, max(len(rows_), 1), len(columns) - 1,
            {
                "data": rows_,
                "columns": [
                    {"header": headers.get(c, c), "header_format": style["head"], "format": formats.get(c, style["text"])}
                    for c in columns
                ],
                "banded_rows": True,
            },
        )

    write_sheet("Emekli-Nüfus oranı", ranked_frame)
    write_sheet("Türkiye", tr)

    notes_sheet(
        book,
        [
            "İl bazında emekli / toplam nüfus oranı — 2025",
            "",
            "Emekli sayısı = malullük + yaşlılık aylığı alanların toplamı (4/1-a +",
            "4/1-b + 4/1-c), ölüm aylığı (dul/yetim) hariç.",
            "",
            "Paydası bilerek TOPLAM nüfus — bir yaş bandına (65+ gibi) bölünmedi.",
            "Sebep: EYT (Emeklilikte Yaşa Takılanlar, 2023) sonrası Türkiye'de",
            "emeklilik yaşı çok geniş bir banda yayıldı, çoğu emekli 45-64 yaşında —",
            "65+ paydası kullanmak emeklilerin büyük bölümünü dışarıda bırakıp oranı",
            "yapay şişirirdi. Toplam nüfus paydası yaş eşleştirmesi iddia etmiyor,",
            "yalnız 'nüfusun ne kadarı emekli maaşı alıyor' sorusuna cevap veriyor.",
            "",
            "Emeklinin fiilen o ilde yaşayıp yaşamadığı ayrı bir soru — SGK kaydı",
            "il müdürlüğüne göre, ikametle birebir örtüşmeyebilir.",
            "",
            "KAYNAK",
            "SGK İstatistik Yıllığı 2025, Tablo 2.5, 2.19, 2.32 (emekli) +",
            "population.csv.gz (TÜİK ADNKS, nüfus).",
        ],
    )
    book.close()
    # endregion

    print("yazildi:", TARGET)


if __name__ == "__main__":
    main()
