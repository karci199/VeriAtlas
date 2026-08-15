"""What each settlement *was* — town, belde or village — from the 2015 election tables.

The 7H workbooks group every settlement into `MAHALLE` (town neighbourhoods) or
`BELDE - KÖY` (the rest, with belde and köy told apart). That grouping is the pre-6360
administrative classification and the YSK kept using it after 2014: Ankara still appears
with 866 villages, İstanbul with 164, in a year when neither had any in law.

That makes this the only reachable source for a question nothing else answers — which of
today's neighbourhoods used to be a village. Kalecik's seven town neighbourhoods here are
exactly the seven a human picked out by hand from the record numbers; Gölköy is a village
here, as it was in the 2000 census.

It is stored as `koken` — origin — and never as the answer to "is this urban". That
question has a tense: Bahçeşehir is a belde in this file and 64.000 people in a dense city
today, and both are true. The density answer lives in `urban_rural` (TÜİK's DEGURBA), the
history lives here, and merging them would destroy the only interesting part.

**Matching is by name**, which everything else in this repo refuses to do — 7H carries no
record number, only İl / İlçe / Mahalle text. The join is therefore made unique-to-unique
inside a district: a name that occurs twice on either side is dropped rather than guessed.
About one settlement in five ends up unmatched and is left `null`. Null, not "koy": a
missing origin is a missing answer, and filling it would put 17 million people into a
category no source put them in.

Run:  uv run python scripts/load_7h.py "<...>/AATOPLU 7H.xlsx"
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import polars as pl

sys.path.insert(0, "src")

from veriatlas.config import RAW

DATA = Path("src/veriatlas/data")
SOURCE_ID = "ysk_7h_2015"

TOWN_SHEET = "MAHALLE"
RURAL_SHEET = "BELDE - KÖY"

#: Trailing words that mean "neighbourhood" or "village" rather than being part of a name.
SUFFIX = re.compile(r"\s*(MAH\.|MAHALLESI|MAH|KOY\.|KOYU)\s*$")
FOLD = str.maketrans("İIŞĞÜÖÇ", "IISGUOC")


def key(text: object) -> str:
    """A name reduced to what two sources can agree on: letters and digits, folded."""
    folded = str(text or "").upper().translate(FOLD)
    return re.sub(r"[^A-Z0-9]", "", SUFFIX.sub("", folded))


def read_7h(path: Path) -> pl.DataFrame:
    """Every settlement in the workbook, as (province, district, name, origin)."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in (TOWN_SHEET, RURAL_SHEET):
        if sheet not in book.sheetnames:
            raise ValueError(f"{path.name} icinde '{sheet}' sayfasi yok")

    rows = []
    # Town sheet: il | ilce | mahalle | ...
    for row in book[TOWN_SHEET].iter_rows(values_only=True):
        if isinstance(row[0], str) and key(row[0]) not in ("IL", "") and row[2]:
            rows.append((key(row[0]), key(row[1]), key(row[2]), "kent"))
    # Rural sheet: il | Belde/Köy | ilce | mahalle | ...
    for row in book[RURAL_SHEET].iter_rows(values_only=True):
        if isinstance(row[0], str) and key(row[0]) not in ("IL", "") and row[3]:
            kind = "belde" if str(row[1] or "").strip() == "Belde" else "koy"
            rows.append((key(row[0]), key(row[2]), key(row[3]), kind))
    if not rows:
        raise ValueError("7H satiri okunamadi — sayfa duzeni degismis olabilir")
    return pl.DataFrame(rows, schema=["il", "ilce", "ad", "koken"], orient="row")


def ours() -> pl.DataFrame:
    """Our settlements with the same three name parts, for both levels."""
    provinces = pl.read_csv(DATA / "areas_tr.csv").select(
        pl.col("area_id").alias("il_id"), pl.col("name_tr").alias("il_adi")
    )
    districts = pl.read_csv(DATA / "areas_tr_districts.csv").select(
        pl.col("area_id").alias("ilce_id"), pl.col("name_tr").alias("ilce_adi")
    )

    def side(name: str, level: str) -> pl.DataFrame:
        return (
            pl.read_csv(DATA / name)
            .with_columns(pl.col("area_id").str.slice(0, 5).alias("il_id"))
            .join(provinces, on="il_id", how="left")
            .join(districts, left_on="parent_id", right_on="ilce_id", how="left")
            .select("area_id", "name_tr", "il_adi", "ilce_adi")
            .with_columns(pl.lit(level).alias("duzey"))
        )

    both = pl.concat(
        [
            side("areas_tr_neighbourhoods.csv", "neighbourhood"),
            side("areas_tr_villages.csv", "village"),
        ]
    )
    return both.with_columns(
        pl.col("il_adi").map_elements(key, return_dtype=pl.Utf8).alias("il"),
        pl.col("ilce_adi").map_elements(key, return_dtype=pl.Utf8).alias("ilce"),
        pl.col("name_tr").map_elements(key, return_dtype=pl.Utf8).alias("ad"),
    )


def matched(theirs: pl.DataFrame, mine: pl.DataFrame) -> pl.DataFrame:
    """Unique-to-unique join on the three name parts.

    `keep="none"` on both sides is the whole safety argument: a district holding two
    "Merkez" keeps neither. Half of a duplicate pair matched at random would be worse than
    no match, because it would look like an answer.
    """
    on = ["il", "ilce", "ad"]
    return mine.unique(subset=on, keep="none").join(
        theirs.unique(subset=on, keep="none"), on=on, how="inner"
    )


def attach(registry: Path, pairs: pl.DataFrame, level: str) -> tuple[int, int]:
    wanted = pairs.filter(pl.col("duzey") == level).select("area_id", "koken")
    frame = pl.read_csv(registry)
    if "koken" in frame.columns:
        frame = frame.drop("koken")
    joined = frame.join(wanted, on="area_id", how="left")
    joined.write_csv(registry)
    return int(joined["koken"].is_not_null().sum()), len(joined)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit('kullanim: load_7h.py "<...>/AATOPLU 7H.xlsx"')
    source = Path(sys.argv[1])
    if not source.exists():
        raise SystemExit(f"dosya yok: {source}")

    RAW.mkdir(parents=True, exist_ok=True)
    kept = RAW / f"ysk-7h-2015-{datetime.now(UTC):%Y-%m-%d}.xlsx"
    if kept.resolve() != source.resolve():
        shutil.copy2(source, kept)
    digest = hashlib.sha256(kept.read_bytes()).hexdigest()[:16]

    theirs = read_7h(kept)
    mine = ours()
    pairs = matched(theirs, mine)
    print(f"7H: {len(theirs)} yerlesim | bizim: {len(mine)} | eslesen: {len(pairs)}")
    print(pairs.group_by("koken").len().sort("koken"))

    report = {}
    for level, registry in (
        ("neighbourhood", DATA / "areas_tr_neighbourhoods.csv"),
        ("village", DATA / "areas_tr_villages.csv"),
    ):
        hit, total = attach(registry, pairs, level)
        report[level] = (hit, total)
        print(f"{level}: {hit}/{total} ({100 * hit / total:.1f}%)")

    manifest = {
        "source_id": SOURCE_ID,
        "dataset": "settlement_origin_7h",
        "file": kept.name,
        "sha256_16": digest,
        "rows": len(theirs),
        "matched": {k: v[0] for k, v in report.items()},
        "of": {k: v[1] for k, v in report.items()},
        "join": "name within district, unique-to-unique",
        "as_of": "2015-06-07",
        "retrieved_at": datetime.now(UTC).strftime("%Y-%m-%d"),
    }
    with (RAW / "manifests.jsonl").open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(manifest, ensure_ascii=False) + "\n")
    print("manifest yazildi:", digest)


if __name__ == "__main__":
    main()
