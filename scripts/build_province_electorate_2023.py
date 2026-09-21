"""Province table for one election: registered voters vs. expected voters.

Expected voters = population aged 18+ (ADNKS single ages) minus foreign nationals aged
18+. Foreign counts by province carry no age split, so the national 18+ share of the
foreign population (TÜİK migration bulletin table, 15-19 band split 3/5) is applied to
every province. Registered voters are domestic ballot boxes summed over each province's
electoral circles.

Adds a "İller <year>" sheet to the Turkey electorate workbook (other sheets untouched).

Usage: python scripts/build_province_electorate_2023.py [year] [workbook]
"""

from __future__ import annotations

import csv
import gzip
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ELECTIONS = Path("C:/veri/raw/tuik_secim_ilce/secim_ilce.csv")
SINGLE_AGE = Path("C:/veri/public/population-age1.csv.gz")
FOREIGN = Path("C:/veri/public/foreign.csv.gz")
FOREIGN_BY_AGE = Path("C:/veri/raw/tuik_foreign_by_age.xls")
DEFAULT_WORKBOOK = Path(r"C:\Users\katan\OneDrive\Desktop\Türkiye Seçmen ve Nüfus.xlsx")

REGISTERED = "Kayıtlı seçmen sayısı"
VOTED = "Oy kullanan seçmen sayısı"
VALID = "Geçerli oy sayısı"
COUNTS = (REGISTERED, VOTED, VALID)
CIRCLE_ALIASES = {"k_maras": "kahramanmaras"}
VOTING_AGE = 18

NAVY, GREY, ZEBRA = "1F3A73", "6B7280", "E9EDF4"
thin = Side(style="thin", color="BFC5D2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def norm(text: str) -> str:
    text = text.replace("İ", "i").replace("I", "ı").lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c)).replace("ı", "i")


def load_elections(year: str) -> dict[str, dict[str, int]]:
    """{normalised province name: counts}; circle total = largest row in that circle."""
    circle_max: dict[str, dict[str, int]] = defaultdict(dict)
    with ELECTIONS.open(encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["yil"] != year or row["olcut"] not in COUNTS:
                continue
            cur = circle_max[row["cevre"]]
            cur[row["olcut"]] = max(cur.get(row["olcut"], 0), int(row["deger"]))
    out: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for circle, counts in circle_max.items():
        province = re.sub(r"_\d+$", "", circle)
        province = CIRCLE_ALIASES.get(province, province)
        for measure, value in counts.items():
            out[norm(province)][measure] += value
    return out


def load_population(year: str):
    """{area_id: (name, total, aged 18+)} for provinces and TR."""
    acc: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    names: dict[str, str] = {}
    with gzip.open(SINGLE_AGE, "rt", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["year"] != year or row["level"] not in ("province", "country"):
                continue
            age = int(row["age"].rstrip("+"))
            value = int(row["value"])
            acc[row["area_id"]]["total"] += value
            if age >= VOTING_AGE:
                acc[row["area_id"]]["adult"] += value
            names[row["area_id"]] = row["area"]
    return {k: (names[k], v["total"], v["adult"]) for k, v in acc.items()}


def load_foreign(year: str) -> dict[str, int]:
    out: dict[str, int] = defaultdict(int)
    with gzip.open(FOREIGN, "rt", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["year"] == year:
                out[row["area_id"]] += int(row["value"])
    return dict(out)


def foreign_adult_share() -> tuple[float, str]:
    """Share of foreign nationals aged 18+ from the national age-group table."""
    sheet = xlrd.open_workbook(FOREIGN_BY_AGE).sheets()[0]
    header = None
    for i in range(sheet.nrows):
        vals = sheet.row_values(i)
        if "0-4" in vals:
            header = {str(v).strip(): j for j, v in enumerate(vals) if v}
        if header and str(vals[1]).startswith("Toplam yabanc"):
            total = vals[3]
            under = (
                sum(vals[header[b]] for b in ("0-4", "5-9", "10-14"))
                + vals[header["15-19"]] * 3 / 5
            )
            ref = str(sheet.row_values(0)[0]).strip()[-4:]
            return 1 - under / total, ref
    raise SystemExit("foreign age table not parsed")


def style(cell, *, header=False, zebra=False, fmt=None, bold=False, align=CENTER):
    if header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    else:
        cell.font = Font(bold=bold, size=10)
        if zebra:
            cell.fill = PatternFill("solid", fgColor=ZEBRA)
    cell.alignment = align
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt


def build_rows(year: str):
    elections = load_elections(year)
    pop = load_population(year)
    foreign = load_foreign(year)
    share, ref_year = foreign_adult_share()
    by_name = {norm(name): area_id for area_id, (name, _, _) in pop.items()}
    rows = []
    for key, counts in elections.items():
        area_id = by_name.get(key)
        if not area_id:
            raise SystemExit(f"province not matched: {key}")
        name, total, adult = pop[area_id]
        rows.append(
            {
                "name": name,
                "registered": counts[REGISTERED],
                "voted": counts[VOTED],
                "total": total,
                "adult": adult,
                "foreign": foreign.get(area_id, 0),
            }
        )
    rows.sort(key=lambda r: r["name"])
    tr_name, tr_total, tr_adult = pop["TR"]
    turkey = {
        "name": tr_name,
        "registered": sum(r["registered"] for r in rows),
        "voted": sum(r["voted"] for r in rows),
        "total": tr_total,
        "adult": tr_adult,
        "foreign": foreign.get("TR", 0),
    }
    return rows, turkey, share, ref_year


def write_sheet(wb, year: str, rows, turkey, share, ref_year):
    title = f"İller {year}"
    if title in wb.sheetnames:
        wb.remove(wb[title])
    ws = wb.create_sheet(title)
    ncol = 12
    for r in range(1, 4):
        for c in range(1, ncol + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=ZEBRA)
    for rng in ("A1:L1", "A2:L2", "A3:L3"):
        ws.merge_cells(rng)
    ws["A1"].value = "SEÇİM · İLLER"
    ws["A1"].font = Font(bold=True, size=10, color=GREY)
    ws[
        "A2"
    ].value = f"Kayıtlı seçmen ve beklenen seçmen — iller, {year} milletvekili seçimi"
    ws["A2"].font = Font(bold=True, size=18, color=NAVY)
    ws["A3"].value = (
        "Beklenen seçmen = 18+ nüfus − 18+ yabancı. Yabancı 18+ payı Türkiye geneli "
        f"(%{share * 100:.1f}, TÜİK {ref_year}) her ile aynı uygulanır. Kayıtlı = yurt içi sandıklar."
    )
    ws["A3"].font = Font(size=10, color=GREY)

    head = [
        "İl",
        "Toplam nüfus",
        "18+ nüfus",
        "Yabancı nüfus",
        "Yabancı 18+",
        "Beklenen seçmen",
        "Kayıtlı seçmen",
        "Beklenen − kayıtlı",
        "Fark / beklenen",
        "Kayıtlı / 18+",
        "Oy kullanan",
        "Katılım",
    ]
    for c, text in enumerate(head, start=1):
        style(ws.cell(5, c, text), header=True)

    def write(r, row, bold=False, zebra=False):
        values = [
            row["name"],
            row["total"],
            row["adult"],
            row["foreign"],
            f"=ROUND(D{r}*{share:.4f},0)",
            f"=C{r}-E{r}",
            row["registered"],
            f"=F{r}-G{r}",
            f"=H{r}/F{r}",
            f"=G{r}/C{r}",
            row["voted"],
            f"=K{r}/G{r}",
        ]
        fmts = [None] + ["#,##0"] * 7 + ["0.0%", "0.0%", "#,##0", "0.0%"]
        for c, (v, f) in enumerate(zip(values, fmts), start=1):
            style(
                ws.cell(r, c, v),
                zebra=zebra,
                fmt=f,
                bold=bold or c == 1,
                align=LEFT if c == 1 else CENTER,
            )

    write(6, turkey, bold=True)
    r = 7
    for i, row in enumerate(rows):
        write(r, row, zebra=i % 2 == 1)
        r += 1

    r += 1
    notes = [
        (
            "Türkiye satırı: 18+ nüfus ADNKS; kayıtlı = illerin toplamı (yurt içi). Yurt dışı ve gümrük seçmeni "
            "ADNKS nüfusunda olmadığı için bu tabloda yok; Türkiye sayfasındaki dipnota bakın."
        ),
        "Eksi fark = ilde kayıtlı seçmen 18+ vatandaş nüfusundan fazla: kütük adresi ilde, fiilen başka yerde yaşayanlar.",
        "Artı fark: er/erbaş, hükümlü, kütük gecikmesi ve yurt dışına gidip kütükten düşmeyenler.",
        "Yabancı nüfus geçici koruma altındaki Suriyelileri kapsamaz (ADNKS dışı); onlar 18+ nüfusta da yok.",
        (
            "Kaynak: TÜİK ADNKS (tek yaş, yabancı nüfus); TÜİK Uluslararası Göç İstatistikleri (yabancı yaş grubu); "
            "TÜİK Seçim Sonuçları Veri Tabanı."
        ),
    ]
    for text in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        cell = ws.cell(r, 1, text)
        cell.font = Font(size=9, color=GREY)
        cell.alignment = LEFT
        ws.row_dimensions[r].height = 26
        r += 1

    for i, w in enumerate([18, 13, 13, 13, 12, 14, 14, 14, 11, 11, 13, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B7"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 130
    ws.auto_filter.ref = f"A5:L{6 + len(rows)}"


def print_summary(rows, turkey, share):
    def gap(r):
        expected = r["adult"] - round(r["foreign"] * share)
        return (
            expected,
            expected - r["registered"],
            (expected - r["registered"]) / expected,
        )

    e, g, p = gap(turkey)
    print(
        f"Türkiye: 18+ {turkey['adult']:,}  yabancı18+ {round(turkey['foreign'] * share):,}  "
        f"beklenen {e:,}  kayıtlı {turkey['registered']:,}  fark {g:,} ({p:.1%})"
    )
    ranked = sorted(rows, key=lambda r: gap(r)[2])
    print("\nEn düşük fark (kütük nüfustan büyük):")
    for r in ranked[:10]:
        e, g, p = gap(r)
        print(f"  {r['name']:<16}{e:>11,}{r['registered']:>11,}{g:>10,}{p:>8.1%}")
    print("\nEn yüksek fark (kütük nüfustan küçük):")
    for r in ranked[-10:][::-1]:
        e, g, p = gap(r)
        print(f"  {r['name']:<16}{e:>11,}{r['registered']:>11,}{g:>10,}{p:>8.1%}")


def main() -> None:
    year = sys.argv[1] if len(sys.argv) > 1 else "2023"
    path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_WORKBOOK
    rows, turkey, share, ref_year = build_rows(year)
    if len(rows) != 81:
        raise SystemExit(f"expected 81 provinces, got {len(rows)}")
    print(f"yabancı 18+ payı: {share:.3%} ({ref_year})")
    print_summary(rows, turkey, share)
    wb = openpyxl.load_workbook(path)
    write_sheet(wb, year, rows, turkey, share, ref_year)
    wb.save(path)
    print(f"\nwritten: {path}")


if __name__ == "__main__":
    main()
