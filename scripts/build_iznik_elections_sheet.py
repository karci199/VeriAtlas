"""Rebuild the "Seçimler" sheet of the İznik workbook from TÜİK district election data.

Only the target sheet is rebuilt; every other sheet (including the user's manual edits)
is left untouched. Source: raw/tuik_secim_ilce/secim_ilce.csv (parliamentary elections,
district level, 1961-2023).

Usage: python scripts/build_iznik_elections_sheet.py [workbook_path]
"""

from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

SOURCE = Path(r"C:\veri\raw\tuik_secim_ilce\secim_ilce.csv")
DEFAULT_WORKBOOK = Path(r"C:\Users\katan\OneDrive\Desktop\İznik.xlsx")
SHEET = "Seçimler"
DISTRICT_TOKEN = "znik"  # matches İznik regardless of dotted-I casing
PROVINCE_PREFIX = "bursa"

NAVY = "1F3A73"
GREY = "6B7280"
ZEBRA = "E9EDF4"
LINK = "0563C1"
ORANGE_FILL = "FDE7C8"

COUNT_KEYS = {
    "Kayıtlı seçmen sayısı": "registered",
    "Oy kullanan seçmen sayısı": "voted",
    "Geçerli oy sayısı": "valid",
    "Sandık say\u0131s\u0131": "ballot_boxes",
}
YEAR_LABELS = {"2015_7_haziran": "2015 Haziran", "2015_1_kasim": "2015 Kasım"}
YEAR_ORDER = {"2015_7_haziran": 2015.0, "2015_1_kasim": 2015.5}

thin = Side(style="thin", color="BFC5D2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def load_district() -> dict[str, dict]:
    """Return {election_key: {counts..., 'parties': {name: votes}}} for the district."""
    elections: dict[str, dict] = defaultdict(lambda: {"parties": {}})
    with SOURCE.open(encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if (
                not row["cevre"].startswith(PROVINCE_PREFIX)
                or DISTRICT_TOKEN not in row["birim"]
            ):
                continue
            rec = elections[row["yil"]]
            value = int(row["deger"])
            key = COUNT_KEYS.get(row["olcut"])
            if key:
                rec[key] = value
            else:
                rec["parties"][row["olcut"]] = value
    return dict(elections)


def sort_key(election: str) -> float:
    return YEAR_ORDER[election] if election in YEAR_ORDER else float(election)


def label(election: str) -> str:
    return YEAR_LABELS.get(election, election)


def style_header(cell, size=None):
    cell.font = Font(bold=True, color="FFFFFF", size=size)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    cell.border = BORDER


def style_body(cell, zebra: bool, *, bold=False, fmt=None, size=10, align=CENTER):
    cell.font = Font(bold=bold, size=size)
    if zebra:
        cell.fill = PatternFill("solid", fgColor=ZEBRA)
    cell.alignment = align
    cell.border = BORDER
    if fmt:
        cell.number_format = fmt


def write_strip(ws, title: str, subtitle: str):
    for r in range(1, 4):
        for c in range(1, 11):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=ZEBRA)
    ws.merge_cells("C1:J1")
    ws.merge_cells("A2:J2")
    ws.merge_cells("A3:J3")
    a1 = ws["A1"]
    a1.value = "« Ana Sayfa"
    a1.hyperlink = Hyperlink(ref="A1", location="'Ana Sayfa'!A1")
    a1.font = Font(bold=True, color=LINK)
    ws["C1"].value = "SEÇİM"
    ws["C1"].font = Font(bold=True, size=10, color=GREY)
    ws["A2"].value = title
    ws["A2"].font = Font(bold=True, size=18, color=NAVY)
    ws["A3"].value = subtitle
    ws["A3"].font = Font(size=10, color=GREY)


def write_section_title(ws, row: int, text: str):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    cell = ws.cell(row, 1, text)
    cell.font = Font(bold=True, size=12, color=NAVY)
    cell.alignment = LEFT


def build(ws, data: dict[str, dict]):
    write_strip(
        ws,
        "Seçimler — milletvekili genel seçimleri 1961–2023",
        "İlçe toplamı, TÜİK seçim dağıtım uygulaması. Üstte katılım ve birinci parti; "
        "altta her seçimdeki parti oyları (sıralanabilir uzun tablo). Turuncu = dipnotlu yıl.",
    )
    elections = sorted(data, key=sort_key)

    # --- Block A: turnout and winner, one row per election -------------------------
    write_section_title(ws, 5, "A. Katılım ve birinci parti")
    head = [
        "Seçim",
        "Kayıtlı seçmen",
        "Oy kullanan",
        "Geçerli oy",
        "Katılım",
        "Kayıp oy",
        "Sandık",
        "Birinci parti",
        "Birinci oy",
        "Birinci pay",
    ]
    for c, text in enumerate(head, start=1):
        style_header(ws.cell(6, c, text))
    row = 7
    footnote_years = {"2007", "2011"}
    year_rows: dict[str, int] = {}
    for i, e in enumerate(elections):
        rec = data[e]
        zebra = i % 2 == 1
        year_rows[e] = row
        parties = {p: v for p, v in rec["parties"].items() if p != "BĞMZ"}
        winner, wvotes = max(parties.items(), key=lambda kv: kv[1])
        values = [
            label(e),
            rec.get("registered"),
            rec.get("voted"),
            rec.get("valid"),
            f"=C{row}/B{row}",
            f"=1-D{row}/B{row}",
            rec.get("ballot_boxes"),
            winner,
            wvotes,
            f"=I{row}/D{row}",
        ]
        fmts = [
            None,
            "#,##0",
            "#,##0",
            "#,##0",
            "0.0%",
            "0.0%",
            "#,##0",
            None,
            "#,##0",
            "0.0%",
        ]
        for c, (v, f) in enumerate(zip(values, fmts), start=1):
            cell = ws.cell(row, c, v)
            style_body(
                cell, zebra, bold=(c == 1), fmt=f, align=LEFT if c == 8 else CENTER
            )
            if e in footnote_years:
                cell.fill = PatternFill("solid", fgColor=ORANGE_FILL)
        row += 1

    # --- Block B: long table of party votes ----------------------------------------
    row += 1
    write_section_title(ws, row, "B. Parti oyları — her seçim, her parti")
    row += 1
    head_b = ["Seçim", "Parti", "Oy", "Pay", "Sıra"]
    for c, text in enumerate(head_b, start=1):
        style_header(ws.cell(row, c, text))
    row += 1
    zebra = False
    for e in elections:
        rec = data[e]
        ranked = sorted(
            ((p, v) for p, v in rec["parties"].items() if v > 0),
            key=lambda kv: -kv[1],
        )
        valid_cell = f"$D${year_rows[e]}"
        for rank, (party, votes) in enumerate(ranked, start=1):
            values = [label(e), party, votes, f"=C{row}/{valid_cell}", rank]
            fmts = [None, None, "#,##0", "0.0%", "0"]
            for c, (v, f) in enumerate(zip(values, fmts), start=1):
                style_body(
                    ws.cell(row, c, v),
                    zebra,
                    bold=(c == 1),
                    fmt=f,
                    align=LEFT if c == 2 else CENTER,
                )
            row += 1
        zebra = not zebra

    # --- Notes ---------------------------------------------------------------------
    row += 1
    notes = [
        "Katılım = oy kullanan / kayıtlı. Kayıp oy = 1 − geçerli / kayıtlı (kullanmayan + geçersiz).",
        "Birinci parti bağımsızlar (BĞMZ) hariç. Pay = geçerli oya bölünür.",
        "2007 ve 2011'de Kürt siyaseti bağımsız adaylarla girdi; BĞMZ o yıllarda parti oyu gibi okunmalı.",
        "1983 seçimine yalnız üç parti (ANAP, HP, MDP) katılabildi. 1980 darbesi sonrası kapatılan partiler yok.",
        "Parti adları TÜİK raporundaki kısaltmalar; blok (sağ / sol / Kürt siyaseti) eşlemesi ayrı sayfada yapılacak.",
        "Kaynak: TÜİK, Seçim Sonuçları Veri Tabanı (biruni.tuik.gov.tr/secimdagitimapp), ilçelere göre tablo.",
    ]
    for text in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1, text)
        cell.font = Font(size=9, color=GREY)
        cell.alignment = LEFT
        ws.row_dimensions[row].height = 18
        row += 1

    widths = [16, 15, 13, 13, 11, 11, 10, 22, 13, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 140


def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    data = load_district()
    if len(data) < 17:
        raise SystemExit(f"expected 17 elections, found {len(data)}: {sorted(data)}")
    wb = openpyxl.load_workbook(path)
    if SHEET in wb.sheetnames:
        index = wb.sheetnames.index(SHEET)
        wb.remove(wb[SHEET])
    else:
        index = len(wb.sheetnames)
    ws = wb.create_sheet(SHEET, index)
    build(ws, data)
    wb.save(path)
    print(f"{SHEET}: {len(data)} elections written to {path}")


if __name__ == "__main__":
    main()
