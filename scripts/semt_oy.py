"""Election results by semt — the postal delivery districts, not the administrative ones.

A semt has no legal definition and no boundary. The only country-wide list is the PTT
postal code table's `semt_bucak_belde` column, and that column is gone from today's PTT
site; the 2022 file is what there is.

Rows where the semt is just the district's own name are dropped. Those are not semts: PTT
uses the district name as a single bucket where it has not divided the district, so
"Sincan / Sincan" is 340.000 people in one box and says nothing, while "Elmadağ /
Hasanoğlan" is a real place with a name of its own.

Neighbourhood populations and vote counts are joined on the area id both sides carry;
the semt is attached by (district, neighbourhood name), folded to an ascii skeleton
because the two sources spell Turkish letters and the "MAH." suffix differently.

Run:  uv run python scripts/semt_oy.py [mv2023] [--il=TR-06] [--n=30] [--min=5000]
"""

from __future__ import annotations

import collections
import json
import os
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
HAM = pathlib.Path(os.environ.get("VERIATLAS_HAM", "C:/veri-ham"))
DEM = HAM / "endeksa" / "demography"
TILES = ROOT / "public" / "tiles"
PTT = HAM / "ptt" / "pk_20220810.xlsx"

PARTILER = [
    ("AK PARTİ", "AK PARTİ"),
    ("CHP", "CHP"),
    ("YEŞİL SOL", "YEŞİL SOL"),
    ("MHP", "MHP"),
    ("İYİ PARTİ", "İYİ PARTİ"),
    ("YENİDEN REFAH", "YENİDEN REFAH"),
    ("ZAFER", "ZAFER"),
]


def fold(text: str | None) -> str:
    """An ascii skeleton, so the two sources' spellings meet."""
    if not text:
        return ""
    text = text.strip().upper()
    for a, b in (("İ", "I"), ("Ş", "S"), ("Ğ", "G"), ("Ü", "U"), ("Ö", "O"), ("Ç", "C")):
        text = text.replace(a, b)
    text = re.sub(r"\b(MAH|MAHALLESI|MH|KOYU|KOY|BLD|BELDESI)\b\.?", "", text)
    return re.sub(r"[^A-Z0-9]", "", text)


def semt_haritasi(province: str) -> dict[tuple[str, str], str]:
    """{(district, neighbourhood): semt} for one province, from the PTT table."""
    wb = openpyxl.load_workbook(PTT, read_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    next(rows)
    out = {}
    for il, ilce, semt, mahalle, _pk in rows:
        if not il or fold(il) != province:
            continue
        out[(fold(ilce), fold(mahalle))] = (semt or "").strip()
    return out


def main(argv: list[str]) -> None:
    vote = next((a for a in argv if not a.startswith("--")), "mv2023")
    il_kodu = next((a.split("=")[1] for a in argv if a.startswith("--il=")), "TR-06")
    n = int(next((a.split("=")[1] for a in argv if a.startswith("--n=")), 30))
    floor = int(next((a.split("=")[1] for a in argv if a.startswith("--min=")), 5000))
    il_adi = next((a.split("=")[1] for a in argv if a.startswith("--iladi=")), "ANKARA")

    oy = {}
    path = TILES / f"secim-{vote}-mahalle-{il_kodu}.json"
    if not path.exists():
        raise SystemExit(f"{path.name} yok — once: parse_secim.py {vote}")
    for area_id, row in json.loads(path.read_text(encoding="utf-8")).items():
        if area_id.count("-") == 3:
            oy[area_id] = row

    semtler = semt_haritasi(fold(il_adi))
    agg = collections.defaultdict(
        lambda: {"nufus": 0.0, "gecerli": 0.0, "kayitli": 0.0, "mah": 0,
                 "v": collections.Counter()}
    )
    disarida = 0
    for dosya in sorted(DEM.glob(f"{il_kodu}-*.json")):
        for ident, record in json.loads(dosya.read_text(encoding="utf-8")).items():
            dem = record.get("demography") or {}
            ilce = dem.get("CountyName") or ""
            semt = semtler.get((fold(ilce), fold(record.get("name_tr"))))
            if not semt or fold(semt) == fold(ilce):
                disarida += 1
                continue  # PTT'nin bölmediği ilçe: semt değil, tek kutu
            row = oy.get(f"{dosya.stem}-{ident}")
            if not row:
                continue
            g = agg[(semt.strip(), ilce)]
            g["nufus"] += dem.get("PopulationTotal") or 0
            g["mah"] += 1
            g["gecerli"] += row.get("g") or sum(row.get("v", {}).values())
            g["kayitli"] += row.get("k") or 0
            for parti, sayi in row.get("v", {}).items():
                g["v"][parti.upper()] += sayi

    rows = [(v["gecerli"], k, v) for k, v in agg.items() if v["gecerli"] >= floor]
    rows.sort(reverse=True)
    print(f"{vote} · {il_adi} · semt: {len(rows)} (gecerli oy >= {floor:,}) · "
          f"ilce adiyla ayni olan ve eslesmeyen: {disarida} mahalle disarida")
    basliklar = " ".join(f"{ad[:9]:>9}" for ad, _ in PARTILER)
    print(f"\n{'semt':22} {'ilce':14} {'gecerli':>9} {'katilim':>7} {basliklar}")
    for gecerli, (semt, ilce), v in rows[:n]:
        katilim = 100 * gecerli / v["kayitli"] if v["kayitli"] else 0
        paylar = " ".join(
            f"{100 * sum(s for p, s in v['v'].items() if parca in p) / gecerli:8.1f}%"
            for _, parca in PARTILER
        )
        print(f"{semt[:22]:22} {ilce[:14]:14} {int(gecerli):9,} {katilim:6.1f}% {paylar}")


if __name__ == "__main__":
    main(sys.argv[1:])
