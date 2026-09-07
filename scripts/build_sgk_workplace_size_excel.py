"""İl başına 4/1-a işyeri büyüklüğü (işyeri başına zorunlu sigortalı) — 2016 vs 2025.

SGK İstatistik Yıllığı 2025, Tablo 1.10: "4/1-a Kapsamında Zorunlu Sigortalıların ve
İş Yerlerinin İle Göre Dağılımı, 2016-2025" — tek tabloda 10 yıllık seri, il başına
hem işyeri sayısı hem zorunlu sigortalı sayısı. Ortalama işyeri büyüklüğü = sigortalı
/ işyeri; büyümesi "aynı işletmeler mi genişliyor yoksa yeni işletme mi açılıyor"
sorusuna cevap veriyor — sigortalı sayısı işyeri sayısından hızlı artıyorsa mevcut
işyerleri büyüyor, ikisi birlikte artıyorsa yeni işyeri açılışı baskın.

Run:  uv run python scripts/build_sgk_workplace_size_excel.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import polars as pl
import xlsxwriter

sys.path.insert(0, "src")
sys.path.insert(0, "scripts")

from build_analysis_books import notes_sheet, ranked, styles

from veriatlas.config import PUBLIC

SCRATCH = Path(
    r"C:\Users\katan\AppData\Local\Temp\claude\C--veri--claude-worktrees-nerede-kaldik-4998d7"
    r"\289f51e5-ac10-4275-b09d-211258ba36af\scratchpad"
)
TARGET = PUBLIC.parent / "cikti" / "analiz-sgk-isyeri-buyuklugu-il-2016-2025.xlsx"

FIRST, LAST = 2016, 2025


def main() -> None:
    path = next((SCRATCH / "sgk2025").glob("*.xlsx"))
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb["TABLO-1.10"]

    rows = []
    for row in ws.iter_rows(min_row=8, max_row=95, values_only=True):
        code, il = row[0], row[1]
        try:
            plate = int(code)
        except (TypeError, ValueError):
            continue
        if not il:
            continue
        isyeri_16 = row[2]
        isyeri_25 = row[11]
        sigortali_16 = row[12]
        sigortali_25 = row[21]
        if isyeri_16 is None or isyeri_25 is None:
            continue
        rows.append(
            {
                "il": str(il).strip(),
                "plaka": plate,
                "isyeri_2016": isyeri_16,
                "isyeri_2025": isyeri_25,
                "sigortali_2016": sigortali_16,
                "sigortali_2025": sigortali_25,
            }
        )
    frame = pl.DataFrame(rows)
    if len(frame) != 81:
        print(f"uyari: 81 il beklenirken {len(frame)} satir okundu")

    frame = frame.with_columns(
        (pl.col("sigortali_2016") / pl.col("isyeri_2016")).alias("buyukluk_2016"),
        (pl.col("sigortali_2025") / pl.col("isyeri_2025")).alias("buyukluk_2025"),
        (pl.col("isyeri_2025") / pl.col("isyeri_2016") - 1).alias("isyeri_oran"),
        (pl.col("sigortali_2025") / pl.col("sigortali_2016") - 1).alias("sigortali_oran"),
    ).with_columns(
        (pl.col("buyukluk_2025") - pl.col("buyukluk_2016")).alias("buyukluk_fark"),
    )

    compare = ranked(
        frame.select(
            "il", "plaka",
            "isyeri_2016", "isyeri_2025", "isyeri_oran",
            "sigortali_2016", "sigortali_2025", "sigortali_oran",
            "buyukluk_2016", "buyukluk_2025", "buyukluk_fark",
        ),
        "buyukluk_fark",
    )

    tr = pl.DataFrame(
        [
            {
                "yil": FIRST,
                "isyeri": frame["isyeri_2016"].sum(),
                "sigortali": frame["sigortali_2016"].sum(),
                "buyukluk": frame["sigortali_2016"].sum() / frame["isyeri_2016"].sum(),
            },
            {
                "yil": LAST,
                "isyeri": frame["isyeri_2025"].sum(),
                "sigortali": frame["sigortali_2025"].sum(),
                "buyukluk": frame["sigortali_2025"].sum() / frame["isyeri_2025"].sum(),
            },
        ]
    )

    # region Yazma
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(str(TARGET))
    style = styles(book)
    left = book.add_format({"align": "left", "valign": "vcenter", "indent": 1})
    size_fmt = book.add_format(
        {"num_format": "0.0", "align": "center", "valign": "vcenter"}
    )

    headers = {
        "il": "İl", "plaka": "Plaka", "sira": "Sıra", "yil": "Yıl",
        "isyeri": "İşyeri sayısı", "sigortali": "Zorunlu sigortalı", "buyukluk": "İşyeri büyüklüğü",
        "isyeri_2016": f"İşyeri · {FIRST}", "isyeri_2025": f"İşyeri · {LAST}", "isyeri_oran": "İşyeri artışı",
        "sigortali_2016": f"Sigortalı · {FIRST}", "sigortali_2025": f"Sigortalı · {LAST}", "sigortali_oran": "Sigortalı artışı",
        "buyukluk_2016": f"İşyeri büyüklüğü · {FIRST}", "buyukluk_2025": f"İşyeri büyüklüğü · {LAST}",
        "buyukluk_fark": "İşyeri büyüklüğü farkı",
    }
    counts = ("isyeri_2016", "isyeri_2025", "sigortali_2016", "sigortali_2025", "isyeri", "sigortali")
    percents = ("isyeri_oran", "sigortali_oran")
    sizes = ("buyukluk_2016", "buyukluk_2025", "buyukluk_fark", "buyukluk")
    formats = {
        "il": left, "plaka": style["text"], "sira": style["text"], "yil": style["text"],
        **{c: style["count"] for c in counts},
        **{c: style["percent"] for c in percents},
        **{c: size_fmt for c in sizes},
    }
    widths = {"il": 22, "plaka": 8, "sira": 7, "yil": 8}

    def write_sheet(name: str, data: pl.DataFrame) -> None:
        page = book.add_worksheet(name)
        page.freeze_panes(1, 1)
        page.set_row(0, 40)
        columns = data.columns
        for index, column in enumerate(columns):
            page.set_column(index, index, widths.get(column, 18), formats.get(column, style["text"]))
        rows_ = [list(r) for r in data.rows()]
        page.add_table(
            0, 0, max(len(rows_), 1), len(columns) - 1,
            {
                "data": rows_,
                "columns": [
                    {"header": headers.get(c, c), "header_format": style["head"], "format": formats.get(c, style["text"])}
                    for c in columns
                ],
                "banded_rows": True,
            },
        )

    write_sheet(f"{FIRST} vs {LAST}", compare)
    write_sheet("Türkiye", tr)

    notes_sheet(
        book,
        [
            f"İl başına 4/1-a işyeri büyüklüğü (işyeri başına zorunlu sigortalı) — {FIRST} vs {LAST}",
            "",
            "İşyeri büyüklüğü = zorunlu sigortalı / işyeri sayısı, yalnız 4/1-a (özel/",
            "kamu işçi statüsü) kapsamında — 4/1-b ve 4/1-c'de 'işyeri' kavramı yok.",
            "",
            "Sigortalı artışı işyeri artışından hızlıysa mevcut işyerleri büyüyor demektir",
            "(istihdam yoğunlaşması); ikisi birlikte artıyorsa yeni işyeri açılışı baskın",
            "(istihdam yayılması) — 'İşyeri büyüklüğü farkı' bu ikisini tek sayıya indirger.",
            "",
            "KAYNAK",
            "SGK İstatistik Yıllığı 2025, Tablo 1.10 (sgk.gov.tr/Istatistik/Yillik) —",
            f"tek tabloda {FIRST}-{LAST} serisi, ayrı yıl dosyası indirmeye gerek kalmadı.",
        ],
    )
    book.close()
    # endregion

    print("yazildi:", TARGET)


if __name__ == "__main__":
    main()
